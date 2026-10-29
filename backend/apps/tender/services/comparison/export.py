"""Экспорт сравнения в xlsx (T5.1; спека этапа 5, §3 T5.1; T-27).

Та же таблица, что ``table.build``: лист «Сравнение» — позиции строками,
по поставщику три колонки (цена, сумма в KZT, примечание: аналог,
комментарий), строки доставки и СМР, итоги и экономия; лист «Поставщики» —
валюта, курс, НДС, срок действия, схема оплаты, покрытие, итоги, доп.
условия. Любой текст участника — через ``htqweb.xlsx_safe.safe_cell``
(формулой не станет).
"""

from __future__ import annotations

import io
from decimal import Decimal

from django.http import HttpResponse
from django.utils.http import content_disposition_header
from openpyxl import Workbook

from apps.tender.services.tenders import tenders
from apps.tender.services.tenders.ctx import Ctx
from htqweb.xlsx_safe import safe_cell

from . import table
from .opening import require_opened

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SCHEMES = {"prepay": "Предоплата 100 %", "advance": "Аванс + остаток", "postpay": "Постоплата"}
EXTRAS = {"delivery": "Доставка", "works": "СМР"}


def _num(value):
    return None if value is None else Decimal(value)


def _note(cell: dict | None) -> str:
    if not cell or not cell.get("participates", True):
        return ""
    parts = []
    if cell.get("is_analog"):
        parts.append(f"аналог: {cell.get('analog_note') or 'без описания'}")
    if cell.get("partial_qty"):
        parts.append(f"кол-во {cell['qty']}")
    if cell.get("included_in_price"):
        parts.append("включена в цену")
    if cell.get("comment") or cell.get("note"):
        parts.append(cell.get("comment") or cell.get("note"))
    return "; ".join(parts)


def _scheme(scheme: dict | None) -> str:
    if not scheme:
        return ""
    text = SCHEMES.get(scheme["kind"], scheme["kind"])
    if scheme["kind"] == "advance":
        first = next((s for s in scheme["stages"] if s["trigger"] == "before_shipment"), None)
        text += f" ({first['pct']} % до отгрузки)" if first else ""
    if scheme.get("postpay_days"):
        text += f", {scheme['postpay_days']} дн."
    return text


def _compare_sheet(book: Workbook, data: dict) -> None:
    sheet = book.create_sheet("Сравнение")
    suppliers = data["suppliers"]
    put = lambda values: sheet.append([safe_cell(sheet, v) for v in values])  # noqa: E731
    put(["№", "Позиция", "Ед.", "Кол-во", "План, KZT",
         *[f"{s['name']}: {title}" for s in suppliers
           for title in ("цена", "сумма, KZT", "примечание")]])
    for item in data["items"]:
        row = [item["line_no"], item["name"], item["uom"]["short_name"], _num(item["qty"]),
               _num(item["plan_amount"])]
        for s in suppliers:
            cell = item["offers"].get(s["invitation_id"]) or {}
            row += [_num(cell.get("price")), _num(cell.get("amount_kzt")), _note(cell)]
        put(row)
    for extra in data["extras"]:
        row = ["", EXTRAS.get(extra["kind"], extra["kind"]), "", None, _num(extra["plan_amount"])]
        for s in suppliers:
            cell = extra["offers"].get(s["invitation_id"])
            row += [_num(cell["amount"]) if cell else None,
                    _num(cell["amount_kzt"]) if cell else None, _note(cell)]
        put(row)
    for title, key in (("Итого по общему набору, KZT", "total_common_kzt"),
                       ("Полная стоимость, KZT", "total_full_kzt"),
                       ("Экономия по общему набору, KZT", "saving_common_kzt")):
        put(["", title, "", None, None,
             *[v for s in suppliers for v in (None, _num(s[key]), "")]])


def _suppliers_sheet(book: Workbook, data: dict) -> None:
    sheet = book.create_sheet("Поставщики")
    put = lambda values: sheet.append([safe_cell(sheet, v) for v in values])  # noqa: E731
    put(["Поставщик", "Валюта", "Курс", "Курс на", "НДС", "Срок действия", "Схема оплаты",
         "Покрытие", "Итого по общему набору, KZT", "Полная стоимость, KZT",
         "Экономия, KZT", "Доп. условия"])
    for s in data["suppliers"]:
        vat = f"с НДС {s['vat_rate']} %" if s["with_vat"] else "без НДС"
        put([s["name"], s["currency_code"], _num(s["rate"]), s["rate_on"], vat,
             s["valid_until"], _scheme(s["payment_scheme"]),
             f"{s['coverage']['offered']} из {s['coverage']['total']}",
             _num(s["total_common_kzt"]), _num(s["total_full_kzt"]),
             _num(s["saving_common_kzt"]), s["extra_terms"]])


def export(ctx: Ctx, tender_id) -> HttpResponse:
    tender = tenders.get_visible(ctx, tender_id)
    require_opened(tender)
    data = table.build(tender)
    book = Workbook(write_only=True)
    _compare_sheet(book, data)
    _suppliers_sheet(book, data)
    buffer = io.BytesIO()
    book.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type=XLSX)
    response["Content-Disposition"] = content_disposition_header(
        True, f"Сравнение {tender.number}.xlsx")
    return response
