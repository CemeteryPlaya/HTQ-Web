"""Экспорт сравнения в xlsx (T5.1; спека этапа 5, §3 T5.1; T-27).

Два листа — «Сравнение» (позиции × поставщики, доставка и СМР, итоги) и
«Поставщики» (условия). Текст участников, похожий на формулу, — только
текстом."""

from __future__ import annotations

import io

import openpyxl
import pytest

from apps.tender.models import OfferLine
from apps.tender.tests import helpers as h
from apps.tender.tests.test_comparison import _opened

pytestmark = pytest.mark.django_db


@pytest.fixture
def world(company_context):
    slug = company_context["slug"]
    h.staff(slug)
    world = _opened(slug)
    from htqweb.tenancy.db import use_company
    with use_company(slug):
        OfferLine.objects.filter(offer__invitation=world.inv["C"]).update(
            comment="=HYPERLINK(\"http://evil\",\"x\")")
    return world


def _book(world, user_id=h.SN):
    resp = h.call("get", f"tenders/{world.tender.id}/comparison/export", world.slug, user_id)
    assert resp.status_code == 200, resp.content
    assert resp["Content-Type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert "attachment" in resp["Content-Disposition"]
    return openpyxl.load_workbook(io.BytesIO(resp.content))


def test_export_has_two_sheets(world):
    book = _book(world)
    assert book.sheetnames == ["Сравнение", "Поставщики"]
    compare = book["Сравнение"]
    flat = [c.value for row in compare.iter_rows() for c in row]
    assert "Швеллер 12П" in flat and "Доставка" in flat and "СМР" in flat
    assert "Итого по общему набору, KZT" in flat
    suppliers = book["Поставщики"]
    names = [row[0].value for row in suppliers.iter_rows(min_row=2)]
    assert names == ["ТОО «Демо Металл»", "ТОО «Демо Электро»", "ТОО «Демо Снаб»"]


def test_export_has_no_formula(world):
    book = _book(world)
    cells = [c for sheet in book.worksheets for row in sheet.iter_rows() for c in row]
    assert not [c.coordinate for c in cells if c.data_type == "f"]
    evil = [c for c in cells if isinstance(c.value, str) and c.value.startswith("=HYPERLINK")]
    assert evil and all(c.data_type == "s" for c in evil)


def test_export_before_opening_is_409(company_context):
    from apps.tender.tests import stage5

    slug = company_context["slug"]
    h.staff(slug)
    world = stage5.build(slug)
    resp = h.call("get", f"tenders/{world.tender.id}/comparison/export", slug, h.FD)
    assert (resp.status_code, resp.json()["code"]) == (409, "E-TND-07")
