"""Разбор книги CashFlow (B6.2, D-B62-3).

Правила чтения — дословно из ``apps/contracts/services/cashflow_import.py`` и
``cashflow_operations_import.py``: книгу ведут руками, и каждое правило
выросло из реальной особенности файла (см. докстринги там). Копия, а не
вызов: ``contracts`` замораживается и когда-нибудь уйдёт, а книгу грузить
дальше. Тест ``test_workbook`` сверяет эту копию со старым разбором на одной
книге — пока старый жив, разойтись им не дадут.

Отличие одно: вид и тип договора — значениями модуля (``works``/``goods``,
``standard``/``open``), а не перечислениями моделей ``contracts``.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import re
from collections import defaultdict
from dataclasses import dataclass
from decimal import ROUND_HALF_UP, Decimal

SHEET_BUDGET = "Бюджет"
SHEET_REGISTRY = "Реестр договоров"
SHEET_OPERATIONS = "Операции"

REGISTRY_COLUMNS = {
    "administrator": "Администратор бюджета",
    "program_code": "Код программы",
    "program_name": "Программа",
    "number": "Номер договора",
    "name": "Наименование договора",
    "counterparty": "Контрагент",
    "bin_iin": "БИН/ИИН",
    "amount": "Сумма договор",
    "advance": "Аванс (ТИП ОПЛАТЫ)",
    "signed_date": "Дата подписания",
    "kind": "Вид",
    "contract_type": "Тип",
    "external_id": "Идентификатор LARK",
}
BUDGET_COLUMNS = {
    "administrator": "Администратор бюджета",
    "program": "Бюджетная программа",
    "limit": "Лимит",
    "currency": "Валюта",
}
# «БИН» в листе «Операции» дважды — нужен второй (полный, секции счёта).
OPERATIONS_COLUMNS = {
    "agreement_lark": ("Идентификатор договора (LARK)", 1),
    "administrator": ("Наименование администратора программы", 1),
    "has_agreement": ("Наличие договора", 1),
    "document_date": ("Дата", 1),
    "amount": ("Сумма счета", 1),
    "bin_iin": ("БИН", 2),
    "supplier": ("Наименование поставщика", 1),
    "goods": ("Наименование ТРУ", 1),
    "budget": ("Бюджет", 1),
}
KIND_BY_LABEL = {"риу": "works", "тмц": "goods"}
TYPE_BY_LABEL = {"стандарт": "standard", "открытый": "open"}
KIND_BY_AGREEMENT = "по договору"
KIND_WITHOUT_AGREEMENT = "без договора"

BIN_LENGTH = 12
MONEY = Decimal("0.01")
SHARE = Decimal("0.001")
ZERO = Decimal("0.00")
EXCEL_EPOCH = dt.date(1899, 12, 30)
DEFAULT_CURRENCY = "KZT"


class WorkbookError(Exception):
    """Книга не той формы — читать нечего, продолжать нельзя."""


@dataclass
class BudgetRow:
    excel_row: int
    administrator: str
    program_code: str
    program_name: str
    limit: Decimal
    currency: str


@dataclass
class RegistryRow:
    excel_row: int
    administrator: str
    program_code: str
    program_name: str
    number: str
    name: str
    counterparty: str
    bin_iin: str
    amount: Decimal
    has_amount: bool
    advance_share: Decimal
    signed_date: dt.date | None
    kind: str
    contract_type: str
    external_id: str


@dataclass
class OperationRow:
    excel_row: int
    kind: str
    agreement_lark: str
    administrator: str
    program_code: str
    program_name: str
    bin_iin: str
    supplier: str
    goods: str
    amount: Decimal
    document_date: dt.date | None
    fingerprint: str = ""


# ── значения ячеек ──────────────────────────────────────────────────────

def text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).replace("\xa0", " ").strip()


def money(value) -> Decimal:
    if value is None or value == "":
        return ZERO
    if isinstance(value, Decimal):
        return value.quantize(MONEY, rounding=ROUND_HALF_UP)
    if isinstance(value, (int, float)):
        return Decimal(repr(value)).quantize(MONEY, rounding=ROUND_HALF_UP)
    cleaned = text(value).replace(" ", "").replace(",", ".")
    return Decimal(cleaned).quantize(MONEY, rounding=ROUND_HALF_UP)


def parse_advance_share(value) -> Decimal:
    if value is None or value == "":
        return Decimal(0)
    raw = text(value)
    is_percent = raw.endswith("%")
    if is_percent:
        raw = raw[:-1].strip()
    raw = raw.replace(",", ".")
    try:
        share = Decimal(raw)
    except Exception as exc:  # noqa: BLE001 — сообщение важнее типа
        raise ValueError(f"не разобран аванс {value!r}") from exc
    if is_percent:
        share = share / Decimal(100)
    if share < 0 or share > 1:
        raise ValueError(f"доля аванса {value!r} вне диапазона 0..1 "
                         f"(проценты пишутся со знаком: «70%», не «70»)")
    return share.quantize(SHARE, rounding=ROUND_HALF_UP)


def parse_excel_date(value) -> dt.date | None:
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, (int, float)):
        return EXCEL_EPOCH + dt.timedelta(days=int(value))
    raw = text(value)
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return dt.datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    if raw.isdigit():
        return EXCEL_EPOCH + dt.timedelta(days=int(raw))
    raise ValueError(f"не разобрана дата подписания {value!r}")


def normalize_bin(raw, declared_length) -> tuple[str, str | None]:
    number = text(raw)
    if not number:
        return "", "БИН пустой"
    warning = None
    declared = text(declared_length)
    if declared.isdigit() and int(declared) != len(number):
        warning = (f"БИН {number}: в колонке длины {declared}, а знаков {len(number)} — "
                   f"формула в книге устарела")
    if len(number) < BIN_LENGTH:
        number = number.rjust(BIN_LENGTH, "0")
    elif len(number) > BIN_LENGTH:
        return number, f"БИН {number} длиннее {BIN_LENGTH} знаков — оставлен как есть"
    return number, warning


def split_program_cell(value) -> tuple[str, str]:
    raw = text(value)
    match = re.match(r"^(\d+)\s+(.+)$", raw)
    if not match:
        return "", raw
    return match.group(1), match.group(2).strip()


def parse_budget_cell(value) -> tuple[str, str]:
    raw = text(value)
    if not raw:
        return "", ""
    if re.fullmatch(r"\d+(\.0+)?", raw):
        return str(int(float(raw))), ""
    return split_program_cell(raw)


# ── листы ────────────────────────────────────────────────────────────────

def _rows(worksheet):
    for index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        yield index, list(row)


def _is_blank(values) -> bool:
    return all(v is None or text(v) == "" for v in values)


def _header_map(header_row, expected: dict[str, str], sheet: str) -> dict[str, int]:
    found = {}
    for index, cell in enumerate(header_row):
        label = text(cell)
        if label:
            found.setdefault(label, index)
    missing = [label for label in expected.values() if label not in found]
    if missing:
        raise WorkbookError(f"лист «{sheet}»: не найдены колонки {missing}. "
                            f"Найдены: {sorted(found)}")
    return {key: found[label] for key, label in expected.items()}


def _sheet(workbook, name):
    if name not in workbook.sheetnames:
        raise WorkbookError(f"в книге нет листа «{name}» (есть: {workbook.sheetnames})")
    return workbook[name]


def read_budget_sheet(workbook) -> list[BudgetRow]:
    rows = _rows(_sheet(workbook, SHEET_BUDGET))
    try:
        _, header = next(rows)
    except StopIteration:
        raise WorkbookError(f"лист «{SHEET_BUDGET}» пуст") from None
    columns = _header_map(header, BUDGET_COLUMNS, SHEET_BUDGET)

    def cell(values, key):
        index = columns[key]
        return values[index] if index < len(values) else None

    result = []
    for excel_row, values in rows:
        if _is_blank(values):
            continue
        code, name = split_program_cell(cell(values, "program"))
        result.append(BudgetRow(
            excel_row=excel_row, administrator=text(cell(values, "administrator")),
            program_code=code, program_name=name, limit=money(cell(values, "limit")),
            currency=text(cell(values, "currency")).upper() or DEFAULT_CURRENCY))
    return result


def read_registry_sheet(workbook) -> tuple[list[RegistryRow], list[str]]:
    rows = _rows(_sheet(workbook, SHEET_REGISTRY))
    try:
        _, header = next(rows)
    except StopIteration:
        raise WorkbookError(f"лист «{SHEET_REGISTRY}» пуст") from None
    columns = _header_map(header, REGISTRY_COLUMNS, SHEET_REGISTRY)
    # Колонка длины БИН без заголовка — сразу за идентификатором LARK.
    bin_length_column = columns["external_id"] + 1
    parsed, warnings = [], []

    def cell(values, index):
        return values[index] if index is not None and index < len(values) else None

    for excel_row, values in rows:
        if _is_blank(values):
            continue
        external_id = text(cell(values, columns["external_id"]))
        number = text(cell(values, columns["number"]))
        if not external_id and not number:
            warnings.append(f"строка {excel_row}: нет ни номера, ни идентификатора — пропущена")
            continue
        try:
            advance_share = parse_advance_share(cell(values, columns["advance"]))
            signed_date = parse_excel_date(cell(values, columns["signed_date"]))
        except ValueError as exc:
            warnings.append(f"строка {excel_row} ({number or external_id}): {exc} — пропущена")
            continue
        bin_iin, bin_warning = normalize_bin(cell(values, columns["bin_iin"]),
                                             cell(values, bin_length_column))
        if bin_warning:
            warnings.append(f"строка {excel_row} ({number or external_id}): {bin_warning}")
        if not bin_iin:
            warnings.append(f"строка {excel_row} ({number or external_id}): без БИН — пропущена")
            continue
        kind_label = text(cell(values, columns["kind"])).lower()
        type_label = text(cell(values, columns["contract_type"])).lower()
        if kind_label and kind_label not in KIND_BY_LABEL:
            warnings.append(f"строка {excel_row} ({number}): вид «{kind_label}» неизвестен — "
                            f"записан как РиУ")
        if type_label and type_label not in TYPE_BY_LABEL:
            warnings.append(f"строка {excel_row} ({number}): тип «{type_label}» неизвестен — "
                            f"записан как стандартный")
        raw_amount = cell(values, columns["amount"])
        has_amount = raw_amount is not None and text(raw_amount) != ""
        parsed.append(RegistryRow(
            excel_row=excel_row, administrator=text(cell(values, columns["administrator"])),
            program_code=text(cell(values, columns["program_code"])),
            program_name=text(cell(values, columns["program_name"])), number=number,
            name=text(cell(values, columns["name"])),
            counterparty=text(cell(values, columns["counterparty"])), bin_iin=bin_iin,
            amount=money(raw_amount) if has_amount else ZERO, has_amount=has_amount,
            advance_share=advance_share, signed_date=signed_date,
            kind=KIND_BY_LABEL.get(kind_label, "works"),
            contract_type=TYPE_BY_LABEL.get(type_label, "standard"),
            external_id=external_id))
    return parsed, warnings


def _header_positions(header_row, expected: dict[str, tuple[str, int]],
                      sheet: str) -> dict[str, int]:
    seen: dict[str, list[int]] = defaultdict(list)
    for index, cell in enumerate(header_row):
        label = text(cell)
        if label:
            seen[label].append(index)
    mapping, missing = {}, []
    for key, (label, occurrence) in expected.items():
        positions = seen.get(label, [])
        if len(positions) >= occurrence:
            mapping[key] = positions[occurrence - 1]
        else:
            missing.append(label if occurrence == 1 else f"{label} (#{occurrence})")
    if missing:
        raise WorkbookError(f"лист «{sheet}»: не найдены колонки {missing}. "
                            f"Найдены: {sorted(seen)}")
    return mapping


def read_operations_sheet(workbook) -> tuple[list[OperationRow], list[str]]:
    """Строки «Операций» с отпечатком (``fingerprint``) — ключом повтора."""
    rows = _rows(_sheet(workbook, SHEET_OPERATIONS))
    marker = OPERATIONS_COLUMNS["has_agreement"][0]
    columns = None
    for _, values in rows:
        if any(text(v) == marker for v in values):
            columns = _header_positions(values, OPERATIONS_COLUMNS, SHEET_OPERATIONS)
            break
    if columns is None:
        raise WorkbookError(f"лист «{SHEET_OPERATIONS}»: не найдена строка заголовков "
                            f"(нет колонки «{marker}»)")

    def cell(values, key):
        index = columns[key]
        return values[index] if index < len(values) else None

    parsed, warnings = [], []
    for excel_row, values in rows:
        kind = text(cell(values, "has_agreement")).lower()
        amount_raw = cell(values, "amount")
        if not kind and text(amount_raw) == "" and not text(cell(values, "supplier")):
            continue
        label = f"строка {excel_row}"
        if kind not in (KIND_BY_AGREEMENT, KIND_WITHOUT_AGREEMENT):
            warnings.append(f"{label}: «Наличие договора» = «{kind}» неизвестно — пропущена")
            continue
        try:
            amount = money(amount_raw)
            document_date = parse_excel_date(cell(values, "document_date"))
        except (ValueError, ArithmeticError) as exc:
            warnings.append(f"{label}: {exc} — пропущена")
            continue
        if amount <= 0:
            warnings.append(f"{label}: сумма счёта пустая или нулевая — пропущена")
            continue
        code, name = parse_budget_cell(cell(values, "budget"))
        bin_iin, _ = normalize_bin(cell(values, "bin_iin"), None)
        parsed.append(OperationRow(
            excel_row=excel_row, kind=kind,
            agreement_lark=text(cell(values, "agreement_lark")),
            administrator=text(cell(values, "administrator")), program_code=code,
            program_name=name, bin_iin=bin_iin, supplier=text(cell(values, "supplier")),
            goods=text(cell(values, "goods")), amount=amount, document_date=document_date))
    seen: dict[tuple, int] = defaultdict(int)
    for row in parsed:
        key = content_key(row)
        row.fingerprint = fingerprint(row, seen[key])
        seen[key] += 1
    return parsed, warnings


def content_key(row: OperationRow) -> tuple:
    return (row.kind, row.administrator, row.agreement_lark, row.bin_iin,
            str(row.amount), row.document_date.isoformat() if row.document_date else "",
            row.goods, row.program_code)


def fingerprint(row: OperationRow, occurrence: int) -> str:
    """``ops:`` + SHA-1 содержимого строки и номера её повтора в книге —
    тот же ключ, что у старого импорта (``Invoice.external_id`` и
    ``ContractPayment.external_id`` в ``contracts``)."""
    raw = "\x1f".join([*map(str, content_key(row)), str(occurrence)])
    return "ops:" + hashlib.sha1(raw.encode("utf-8")).hexdigest()


def load(path):
    """Книга целиком: ``(бюджет, реестр, операции, предупреждения)``.
    Листа «Операции» может не быть — тогда операций нет."""
    from openpyxl import load_workbook

    try:
        book = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — любой отказ чтения — «не та книга»
        raise WorkbookError(f"книга не открывается: {exc}") from exc
    budget = read_budget_sheet(book)
    registry, warnings = read_registry_sheet(book)
    operations: list[OperationRow] = []
    if SHEET_OPERATIONS in book.sheetnames:
        operations, op_warnings = read_operations_sheet(book)
        warnings += op_warnings
    return budget, registry, operations, warnings
