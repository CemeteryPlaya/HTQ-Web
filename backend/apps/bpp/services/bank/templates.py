"""Чтение выписки Excel и CSV по шаблону (ТЗ §11.3 п.1, задача A3.1; Review
Focus 3). Этим же модулем пользуется загрузка выписки (задача 3).

Колонки ищутся по заголовку, а не по номеру: строка заголовка — первая из
первых ``HEADER_SCAN_ROWS`` строк, где нашлись все обязательные поля шаблона
(у банков над таблицей бывает шапка «Выписка по счёту … за период …»).
Регистр, пробелы и «ё» в заголовке не важны, лишние колонки не мешают.
Обязательной колонки нет ни в одной строке — ``E-IMP-02`` с её названием.

Деньги — ``Decimal``, никогда не ``float``: числовая ячейка xlsx приходит из
openpyxl как ``float`` и переводится через ``repr`` (кратчайшая запись того
же числа), текст «1 250 000,00» — разбором строки (``parse_amount``).

Строка с нераспознанной датой или суммой уходит в список ошибок «Строка 17:
не распознана дата „31.02.2026“» (текст ТЗ §11.3 п.1), остальные строки
читаются дальше. Сумма, которую не вместит столбец строки выписки (больше
``AMOUNT_MAX_DIGITS`` цифр до запятой), — тоже ошибка строки, а не отказ
всей выписки при записи.

Выписка 1С сюда не относится: её поля задаёт стандарт формата, разбор —
``parsers/onec.py`` (задача 3). Предпросмотр шаблона 1С поэтому отвечает
422: показывать в нём нечего — колонок у шаблона нет (решение задачи 2).
"""

from __future__ import annotations

import codecs
import csv
import io
import re
import zipfile
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from typing import Iterable, Iterator
from xml.etree.ElementTree import ParseError

from defusedxml import DefusedXmlException

from apps.bpp.models.bank import AmountMode, StatementFormat
from htqweb.errors import DomainError

__all__ = [
    "AMOUNT_FIELDS",
    "AMOUNT_MAX_DIGITS",
    "AmountTooLarge",
    "EXTENSIONS",
    "FIELDS",
    "HEADER_SCAN_ROWS",
    "PREVIEW_ROWS",
    "amount_error",
    "check_extension",
    "date_pattern",
    "find_header",
    "format_error",
    "normalize_header",
    "parse_amount",
    "parse_date",
    "preview",
    "read_row",
    "required_fields",
    "resolve_columns",
    "table_rows",
]

#: Поле выписки → название для человека (подпись в форме шаблона, текст
#: ошибки E-IMP-02, колонка предпросмотра).
FIELDS: dict[str, str] = {
    "date": "Дата",
    "doc_number": "Номер документа",
    "amount": "Сумма",
    "debit": "Дебет",
    "credit": "Кредит",
    "currency": "Валюта",
    "payer_account": "Счёт плательщика",
    "recipient_name": "Получатель",
    "recipient_bin": "БИН/ИИН получателя",
    "recipient_iban": "IBAN получателя",
    "purpose": "Назначение платежа",
}
AMOUNT_FIELDS = {AmountMode.SIGNED: ("amount",), AmountMode.SPLIT: ("debit", "credit")}

#: Допустимые расширения файла по формату шаблона (ТЗ §11.2: «Расширение
#: соответствует формату»). CSV банк-клиенты нередко отдают с ``.txt``.
EXTENSIONS = {
    StatementFormat.ONEC: (".txt",),
    StatementFormat.XLSX: (".xlsx",),
    StatementFormat.CSV: (".csv", ".txt"),
}

HEADER_SCAN_ROWS = 30
PREVIEW_ROWS = 20
CENT = Decimal("0.01")

#: Потолок распакованного xlsx. Файл до 20 МБ сжимается в десятки раз, и
#: «zip-бомба» развернулась бы в памяти воркера: объявленные размеры частей
#: сверяются до открытия книги. Настоящая выписка на порядки меньше.
XLSX_UNPACKED_MAX_MB = 200

#: Короткое имя формата для текста E-IMP-01.
_FORMAT_NAMES = {StatementFormat.ONEC: "1С", StatementFormat.XLSX: "Excel",
                 StatementFormat.CSV: "CSV"}

#: Что бросает openpyxl на битой книге: оборванный XML листа (``ParseError``
#: — при чтении строк, не при открытии), нет нужной части архива
#: (``KeyError``), книга без листов (``IndexError``), сущности XML, которые
#: запрещает defusedxml, и прочее «это не xlsx».
_XLSX_ERRORS = (zipfile.BadZipFile, ParseError, DefusedXmlException, IndexError, KeyError,
                ValueError, OSError)


# ── поля шаблона ────────────────────────────────────────────────────────

def required_fields(amount_mode: str) -> tuple[str, ...]:
    """Обязательные поля: дата, номер документа, сумма (одна колонка или
    дебет и кредит — по режиму), назначение."""
    return ("date", "doc_number", *AMOUNT_FIELDS[AmountMode(amount_mode)], "purpose")


def normalize_header(value) -> str:
    """Заголовок для сравнения: без регистра, без пробелов (включая
    неразрывные и переводы строк внутри ячейки), «ё» = «е»."""
    if value is None:
        return ""
    return re.sub(r"\s+", "", str(value)).casefold().replace("ё", "е")


# ── даты ───────────────────────────────────────────────────────────────

_DATE_TOKENS = (("ГГГГ", "%Y"), ("YYYY", "%Y"), ("ГГ", "%y"), ("YY", "%y"),
                ("ММ", "%m"), ("MM", "%m"), ("ДД", "%d"), ("DD", "%d"))


def date_pattern(mask: str) -> str:
    """Маска «ДД.ММ.ГГГГ» (или латиницей «DD.MM.YYYY») → шаблон ``strptime``.
    Без дня, месяца или года — ``ValueError``."""
    pattern = (mask or "").strip().upper()
    for token, directive in _DATE_TOKENS:
        pattern = pattern.replace(token, directive)
    for directive in ("%d", "%m"):
        if pattern.count(directive) != 1:
            raise ValueError(mask)
    if pattern.count("%Y") + pattern.count("%y") != 1:
        raise ValueError(mask)
    return pattern


def parse_date(value, pattern: str) -> date:
    """Дата ячейки: ``datetime``/``date`` из xlsx — как есть, текст — по
    шаблону; хвост времени («05.09.2026 10:15:00») отбрасывается.
    Не дата — ``ValueError``."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    if not text:
        raise ValueError(value)
    try:
        return datetime.strptime(text, pattern).date()
    except ValueError:
        head = text.split()[0]
        if head == text:
            raise
        return datetime.strptime(head, pattern).date()


# ── суммы ──────────────────────────────────────────────────────────────

#: Цифр в целой части суммы — не больше: столбец ``amount`` строки выписки
#: — ``numeric(18, 2)``. Сумма крупнее (числовая ячейка xlsx «1E+20»,
#: опечатка в выгрузке) переполнила бы его при записи, и отказом стала бы
#: вся выписка с «внутренней ошибкой», а не одна строка.
AMOUNT_MAX_DIGITS = 16
_AMOUNT_LIMIT = Decimal(10) ** AMOUNT_MAX_DIGITS


class AmountTooLarge(ValueError):
    """Сумма распознана, но в ней больше ``AMOUNT_MAX_DIGITS`` цифр до
    запятой — ошибка строки со своим текстом (``amount_error``)."""


def amount_error(row_no: int, shown: str, exc: ValueError | None = None) -> str:
    """Текст ошибки строки о сумме «Строка N: …»: слишком большая
    (``AmountTooLarge``) или не распознана."""
    if isinstance(exc, AmountTooLarge):
        return (f"Строка {row_no}: сумма слишком большая „{shown}“ — в сумме не больше "
                f"{AMOUNT_MAX_DIGITS} цифр до запятой")
    return f"Строка {row_no}: не распознана сумма „{shown}“"


_SPACES = re.compile(r"[\s   '’]+")
_NUMBER = re.compile(r"-?\d+(?:\.\d+)?")


def parse_amount(value) -> Decimal | None:
    """Сумма ячейки → ``Decimal`` с двумя знаками (ROUND_HALF_UP); пусто —
    ``None``, не число — ``ValueError``.

    ``float`` (числовая ячейка xlsx) — через ``repr``: ``1250000.1`` даёт
    ``Decimal("1250000.10")``, а не хвост двоичной дроби. Текст: пробелы
    (и неразрывные) и апострофы — разделители разрядов; из запятой и точки
    десятичный — тот, что правее; одна запятая без точки — десятичная
    («1 250 000,00»), несколько одинаковых — разряды («1,250,000»).

    Поэтому «1,250» — это 1,25, а не тысяча двести пятьдесят: намеренно.
    Банки Казахстана пишут дробную часть через запятую, а разряды —
    пробелом; американская запись «1,250» без точки в их выгрузках не
    встречается, и угадывать по числу цифр после запятой опаснее, чем
    держаться местного правила.

    Больше ``AMOUNT_MAX_DIGITS`` цифр до запятой (после округления) —
    ``AmountTooLarge``: такую сумму не вместит столбец строки выписки.
    """
    if value is None:
        return None
    if isinstance(value, bool):
        raise ValueError(value)
    if isinstance(value, Decimal):
        number = value
    elif isinstance(value, int):
        number = Decimal(value)
    elif isinstance(value, float):
        number = Decimal(repr(value))
    else:
        text = _SPACES.sub("", str(value)).replace("−", "-")
        if not text:
            return None
        if text.startswith("(") and text.endswith(")"):  # бухгалтерский минус
            text = "-" + text[1:-1]
        comma, dot = text.rfind(","), text.rfind(".")
        if comma >= 0 and dot >= 0:
            thousands = "." if comma > dot else ","
            text = text.replace(thousands, "")
            text = text.replace(",", ".")
        elif comma >= 0:
            text = text.replace(",", ".") if text.count(",") == 1 else text.replace(",", "")
        elif text.count(".") > 1:
            text = text.replace(".", "")
        if not _NUMBER.fullmatch(text):
            raise ValueError(value)
        number = Decimal(text)
    # До округления — иначе «1E+30» упал бы в ``quantize`` как не число
    # (точности контекста не хватает на 33 знака), а не как большая сумма.
    if number.is_finite() and abs(number) >= _AMOUNT_LIMIT:
        raise AmountTooLarge(value)
    try:
        number = number.quantize(CENT, rounding=ROUND_HALF_UP)
    except InvalidOperation as exc:  # nan, бесконечность
        raise ValueError(value) from exc
    if abs(number) >= _AMOUNT_LIMIT:  # 9999999999999999,995 округлилось вверх
        raise AmountTooLarge(value)
    return number


# ── формат файла ────────────────────────────────────────────────────────

def format_error(name: str, template) -> DomainError:
    """Файл не того формата, что у шаблона (E-IMP-01; ТЗ §26.1 — «Выберите
    другой формат или файл»). Имя файла в текст не входит — его показывает
    форма рядом с полем; ``name`` оставлен в сигнатуре ради вызывающих."""
    fmt = StatementFormat(template.format)
    allowed = " или ".join(EXTENSIONS[fmt])
    message = (f"Файл не соответствует формату {_FORMAT_NAMES[fmt]}: нужен файл {allowed}. "
               f"Выберите другой формат или файл.")
    return DomainError("E-IMP-01", message, fields=[{"field": "file", "message": message}])


def check_extension(name: str, template) -> None:
    allowed = EXTENSIONS[StatementFormat(template.format)]
    if not (name or "").lower().endswith(allowed):
        raise format_error(name, template)


# ── строки таблицы ─────────────────────────────────────────────────────

def _check_unpacked_size(content: bytes, name: str, template) -> None:
    """Сумма объявленных размеров частей архива — до распаковки (zip-бомба).
    Не zip вовсе — E-IMP-01, как и у ``load_workbook``."""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            unpacked = sum(info.file_size for info in archive.infolist())
    except (zipfile.BadZipFile, OSError, ValueError) as exc:
        raise format_error(name, template) from exc
    if unpacked > XLSX_UNPACKED_MAX_MB * 1024 * 1024:
        message = (f"Файл «{name}» после распаковки больше {XLSX_UNPACKED_MAX_MB} МБ — "
                   f"это не похоже на выписку. Выберите другой формат или файл.")
        raise DomainError("E-IMP-01", message, fields=[{"field": "file", "message": message}])


def _xlsx_rows(content: bytes, name: str, template) -> Iterator[tuple]:
    import openpyxl  # тяжёлый импорт — только когда выписку реально читают

    _check_unpacked_size(content, name, template)
    try:
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except _XLSX_ERRORS as exc:
        raise format_error(name, template) from exc
    try:
        # Битый лист обнаруживается не при открытии, а на чтении строк —
        # поэтому под перехватом и чтение: иначе вместо E-IMP-01 был бы 500.
        try:
            sheet = book.worksheets[0]
            # Размер листа из файла у выгрузок банк-клиентов бывает ложным
            # («A1»), и read_only прочитал бы одну строку: считаем заново.
            sheet.reset_dimensions()
            yield from sheet.iter_rows(values_only=True)
        except _XLSX_ERRORS as exc:
            raise format_error(name, template) from exc
    finally:
        book.close()


def _csv_rows(content: bytes, name: str, template) -> Iterator[list]:
    encoding = (template.encoding or "cp1251").strip()
    if (content.startswith(codecs.BOM_UTF8)
            or encoding.lower().replace("-", "").replace("_", "") in ("utf8", "utf8sig")):
        # BOM — признак UTF-8 сильнее кодировки шаблона: Excel сохраняет CSV
        # «UTF-8 с BOM» и у банка с шаблоном cp1251, а в cp1251 BOM прочитался
        # бы как «п»ї» и прилип к первому заголовку.
        encoding = "utf-8-sig"
    try:
        text = content.decode(encoding)
    except (UnicodeDecodeError, LookupError) as exc:
        message = (f"Файл «{name}» не читается в кодировке {template.encoding} шаблона "
                   f"«{template.name}». Проверьте файл или кодировку в шаблоне.")
        raise DomainError("E-IMP-01", message,
                          fields=[{"field": "file", "message": message}]) from exc
    try:
        yield from csv.reader(io.StringIO(text, newline=""),
                              delimiter=template.delimiter or ";")
    except csv.Error as exc:  # NUL-байты, оборванные кавычки — это не CSV
        raise format_error(name, template) from exc


def table_rows(content: bytes, name: str, template) -> Iterator[tuple[int, list]]:
    """Строки таблицы файла с номерами строк файла (с 1). Только Excel и
    CSV; формат проверяется по расширению (E-IMP-01)."""
    check_extension(name, template)
    if template.format == StatementFormat.XLSX:
        rows = _xlsx_rows(content, name, template)
    elif template.format == StatementFormat.CSV:
        rows = _csv_rows(content, name, template)
    else:
        raise ValueError(f"формат {template.format} не табличный")
    try:
        for row_no, row in enumerate(rows, start=1):
            yield row_no, list(row)
    finally:
        rows.close()  # книга xlsx закрывается и когда файл не дочитан


# ── заголовок ──────────────────────────────────────────────────────────

def resolve_columns(header_row: Iterable, template) -> dict[str, int]:
    """Поля шаблона, найденные в строке ``header_row``: поле → номер колонки
    (с 0). Сравнение — ``normalize_header``; из повторившихся заголовков
    берётся первый. Не найденных полей в ответе нет."""
    wanted = {normalize_header(text): field for field, text in (template.columns or {}).items()
              if normalize_header(text)}
    found: dict[str, int] = {}
    for index, cell in enumerate(header_row):
        field = wanted.get(normalize_header(cell))
        if field and field not in found:
            found[field] = index
    return found


def _missing_error(missing: list[str], template) -> DomainError:
    names = ", ".join(f"«{template.columns.get(f) or FIELDS[f]}»" for f in missing)
    if len(missing) == 1:
        text = f"Не найдена обязательная колонка {names}"
    else:
        text = f"Не найдены обязательные колонки {names}"
    message = (f"{text} шаблона «{template.name}» в первых {HEADER_SCAN_ROWS} строках "
               f"файла. Проверьте, что выписка выгружена в формате этого шаблона, "
               f"или поправьте заголовки в шаблоне.")
    return DomainError("E-IMP-02", message, fields=[
        {"field": field, "message": template.columns.get(field) or FIELDS[field]}
        for field in missing])


def find_header(rows: Iterator[tuple[int, list]], template) -> tuple[int, dict[str, int]]:
    """Найти строку заголовка среди первых ``HEADER_SCAN_ROWS`` строк
    ``rows`` (итератор ``table_rows``, дальше его читают с места после
    заголовка). Возвращает (номер строки, поле → колонка). Нет строки со
    всеми обязательными полями — ``E-IMP-02`` с недостающими колонками
    самой полной из строк-кандидатов."""
    required = required_fields(template.amount_mode)
    best: list[str] = list(required)
    for row_no, row in rows:
        found = resolve_columns(row, template)
        missing = [field for field in required if field not in found]
        if not missing:
            return row_no, found
        if len(missing) < len(best):
            best = missing
        if row_no >= HEADER_SCAN_ROWS:
            break
    raise _missing_error(best, template)


# ── строка данных ──────────────────────────────────────────────────────

def _blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _text(value) -> str:
    """Текст ячейки: число без хвоста «.0» (номер документа, БИН в числовой
    ячейке xlsx), остальное — строкой без крайних пробелов."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.date().isoformat()
    return str(value).strip()


def _quote(value) -> str:
    return _text(value) if not isinstance(value, float) else repr(value)


def _bin_text(value) -> str:
    """БИН/ИИН — 12 цифр. В числовой ячейке xlsx ведущий ноль теряется
    (050140000656 → 50140000656): число дополняется нулями слева до 12."""
    if isinstance(value, bool):
        return _text(value)
    if isinstance(value, int) or (isinstance(value, float) and value.is_integer()):
        return str(int(value)).zfill(12)
    return _text(value)


def read_row(row: list, columns: dict[str, int], template, pattern: str,
             row_no: int) -> dict | None:
    """Строка данных → словарь полей выписки; пустая строка — ``None``;
    нераспознанная — ``ValueError`` с текстом ошибки «Строка N: …».

    Строка без даты, номера документа и суммы разом тоже ``None``: это
    подвал или шапка банка между строками («Остаток на конец дня»,
    подпись), а не операция. Строка «Итого» с суммой, но без номера
    документа — уже ошибка строки: деньги в ней есть, и молча терять их
    нельзя.

    ``amount`` — сумма по модулю, ``direction`` — ``debit`` (списание) или
    ``credit`` (поступление): в режиме ``signed`` списание — отрицательная
    сумма, в ``split`` — непустой дебет (знак в колонке не важен: часть
    банков пишет дебет с минусом). Сумма и в дебете, и в кредите, нулевая
    сумма — ошибка строки."""

    def cell(field):
        index = columns.get(field)
        return row[index] if index is not None and index < len(row) else None

    if all(_blank(value) for value in row):
        return None
    key_fields = ("date", "doc_number", *AMOUNT_FIELDS[AmountMode(template.amount_mode)])
    if all(_blank(cell(field)) for field in key_fields):
        return None
    raw_date = cell("date")
    try:
        doc_date = parse_date(raw_date, pattern)
    except ValueError:
        raise ValueError(f"Строка {row_no}: не распознана дата „{_quote(raw_date)}“") from None

    if template.amount_mode == AmountMode.SPLIT:
        amounts = {}
        for field in ("debit", "credit"):
            try:
                amounts[field] = abs(parse_amount(cell(field)) or Decimal("0.00"))
            except ValueError as exc:
                raise ValueError(amount_error(row_no, _quote(cell(field)), exc)) from None
        if amounts["debit"] and amounts["credit"]:
            raise ValueError(f"Строка {row_no}: сумма и в дебете, и в кредите — "
                             f"у операции должна быть одна")
        if amounts["debit"]:
            direction, amount = "debit", amounts["debit"]
        elif amounts["credit"]:
            direction, amount = "credit", amounts["credit"]
        else:
            raise ValueError(f"Строка {row_no}: нет суммы ни в дебете, ни в кредите")
    else:
        raw_amount = cell("amount")
        try:
            signed = parse_amount(raw_amount)
        except ValueError as exc:
            raise ValueError(amount_error(row_no, _quote(raw_amount), exc)) from None
        if signed is None:
            raise ValueError(amount_error(row_no, _quote(raw_amount)))
        if not signed:
            raise ValueError(f"Строка {row_no}: нулевая сумма")
        direction = "debit" if signed < 0 else "credit"
        amount = abs(signed)

    doc_number = _text(cell("doc_number"))
    if not doc_number:
        raise ValueError(f"Строка {row_no}: нет номера документа")
    values = {"row_no": row_no, "date": doc_date, "doc_number": doc_number,
              "amount": amount, "direction": direction}
    for field in ("currency", "payer_account", "recipient_name", "recipient_iban", "purpose"):
        values[field] = _text(cell(field))
    values["recipient_bin"] = _bin_text(cell("recipient_bin"))
    return values


# ── предпросмотр ───────────────────────────────────────────────────────

def _read(upload) -> tuple[bytes, str]:
    name = getattr(upload, "name", "") or ""
    if hasattr(upload, "seek"):
        upload.seek(0)
    return upload.read(), name


def preview(upload, template) -> dict:
    """Предпросмотр образца выписки по шаблону — без сохранения чего-либо.

    ``{header_row, columns: [{field, label, header, index}], rows: [первые
    PREVIEW_ROWS строк данных], errors: ["Строка N: …"]}``. ``index`` —
    номер колонки в файле (с 0), ``None`` — колонка шаблона не нашлась
    (необязательная; нет обязательной — ``E-IMP-02``). В счёт
    ``PREVIEW_ROWS`` идут и строки с ошибкой: предпросмотр показывает
    начало файла, а не первые 20 удачных строк."""
    if template.format == StatementFormat.ONEC:
        message = ("Предпросмотр доступен для шаблонов Excel и CSV. Выписка 1С "
                   "(1CClientBankExchange) разбирается по стандарту формата — колонки "
                   "в шаблоне ей не нужны.")
        raise DomainError("E-VAL-01", message, fields=[{"field": "format", "message": message}])
    content, name = _read(upload)
    pattern = date_pattern(template.date_format)
    rows = table_rows(content, name, template)
    result_rows, errors, seen = [], [], 0
    try:
        header_row, columns = find_header(rows, template)
        for row_no, row in rows:
            try:
                values = read_row(row, columns, template, pattern, row_no)
            except ValueError as exc:
                errors.append(str(exc))
                seen += 1
            else:
                if values is None:
                    continue
                result_rows.append(values)
                seen += 1
            if seen >= PREVIEW_ROWS:
                break
    finally:
        rows.close()  # дальше PREVIEW_ROWS строк файл не читается
    return {
        "header_row": header_row,
        "columns": [{"field": field, "label": FIELDS.get(field, field), "header": header,
                     "index": columns.get(field)}
                    for field, header in (template.columns or {}).items()],
        "rows": result_rows,
        "errors": errors,
    }
