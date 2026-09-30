"""Сальдо и сверка переноса (B6.1, задача 6, D-B61-1): закрытое остаётся в
«Договорах», но его вклад держит сальдо; остатки сходятся с точностью до
ожидаемых расхождений, иначе — стоп и откат."""

from __future__ import annotations

from decimal import Decimal

import pytest
from django.core.management import CommandError
from openpyxl import load_workbook

from apps.bpp.models import Agreement, PurchaseRequest
from apps.bpp.services.budget import committed
from apps.bpp.services.migration import reconcile
from apps.bpp.tests.test_files import memory_storage  # noqa: F401  (фикстура)
from apps.project import interface as projects

from .common import migrate, write_maps
from .test_documents import _world

pytestmark = pytest.mark.django_db


def test_saldo_keeps_closed_spending_and_remaining_matches(company_context, tmp_path):
    slug = company_context["slug"]
    admin, steel, *_ = _world()
    paths = write_maps(tmp_path, {admin.id: "ОБ-А"}, {steel.id: "T-METAL"})
    migrate(slug, *paths, "--report", str(tmp_path / "report.xlsx"))

    # «Занято» в «Договорах» — 685 000; перенесённые документы по старым
    # правилам — 622 000; сальдо = исполненный договор 50 000 + оплаченный
    # счёт 5 000 + закрытый подотчёт 8 000.
    saldo = PurchaseRequest.objects.get(justification__startswith="Сальдо до перехода")
    assert saldo.items.get().amount == Decimal("63000.00") and saldo.is_migrated

    # Единственное ожидаемое расхождение — счёт «На согласовании», ставший
    # черновиком с резервом технической заявки: 9 000.
    project_id = projects.project_ids_by_code(["ОБ-А"])["ОБ-А"]
    article_id = str(Agreement.objects.first().article_id)
    assert committed.committed_by_article(project_id)[article_id] == Decimal("694000.00")

    book = load_workbook(tmp_path / "report.xlsx")
    header, *rows = list(book["Сверка"].values)
    row = dict(zip(header, rows[0]))
    assert (row["old_committed"], row["new_committed"], row["expected_diff"], row["diff"]) == (
        685000, 694000, 9000, 0)
    header, *details = list(book["Ожидаемые расхождения"].values)
    assert [(dict(zip(header, detail))["kind"], dict(zip(header, detail))["delta"])
            for detail in details] == [("счёт", 9000)]
    assert {"Переотправить", "Документы", "Не перенесено"} <= set(book.sheetnames)


def test_unexpected_difference_stops_and_rolls_back(company_context, tmp_path, monkeypatch):
    slug = company_context["slug"]
    admin, steel, *_ = _world()
    paths = write_maps(tmp_path, {admin.id: "ОБ-А"}, {steel.id: "T-METAL"})
    real = reconcile.calc.committed_by_article

    def skewed(project_id, article_ids=None, **kwargs):
        return {key: value + 1 for key, value in real(project_id, article_ids, **kwargs).items()}

    monkeypatch.setattr(reconcile.calc, "committed_by_article", skewed)
    with pytest.raises(CommandError) as exc:
        migrate(slug, *paths)
    assert "остатки не сошлись" in str(exc.value) and "лишнее 1" in str(exc.value)
    assert not projects.project_ids_by_code(["ОБ-А"])
    assert not Agreement.objects.exists()
