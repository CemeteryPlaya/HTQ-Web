"""Ручные действия ФД в сверке выписки (ТЗ §11.2–11.4, BR-060, BR-073; план
этапа 4 A, задача 3 — Review Focus 3 и 5): кандидаты, ручное
сопоставление, подтверждение «Требует проверки», отмена сопоставления,
исключение строки, отмена загрузки, «Сверить», выгрузка результата, права
и идемпотентность.

Счета, где поток B не проверяется, заводятся прямо в базе
(``common.orm_invoice``); где нужна отметка оплаты БУХ — фабриками B.
"""

from __future__ import annotations

import io
from decimal import Decimal

import openpyxl
import pytest
from django.db.models import Sum
from django.test import Client
from django.utils import timezone

from apps.bpp.models import AuditLog, Invoice
from apps.bpp.models.bank import (
    BankImportStatus,
    BankStatementLine,
    LineMatchStatus,
    PaymentMatch,
    PaymentMatchState,
)
from apps.bpp.models.invoices import ReconStatus
from apps.bpp.services.bank import imports, matching, recon
from apps.bpp.services.invoices import payments
from apps.bpp.tests import stage2 as s
from apps.bpp.tests import test_invoices as invoice_flow
from apps.bpp.tests.test_files import memory_storage  # noqa: F401  (фикстура)
from htqweb.errors import DomainError
from htqweb.tenancy.db import use_company

from . import common
from .test_matching import CP_BIN, FOREIGN_BIN, _doc, _payable, _recon, _upload

D = Decimal
FD, BUH, TD = s.FD, 907, s.TD
BASE = "/api/bpp/v1/bank"
SHORT = "123456789"                 # 9 символов — BR-060
REASON = "Проверено по платёжке банка"


@pytest.fixture
def slug(company_context):
    slug = company_context["slug"]
    s.grant(slug, FD, "bpp-fd")
    s.grant(slug, BUH, "bpp-buh")
    s.grant(slug, TD, "bpp-td")
    return slug


def _error(fn, *args, **kwargs) -> DomainError:
    with pytest.raises(DomainError) as exc:
        fn(*args, **kwargs)
    return exc.value


def _line(amount="1000", purpose="Оплата за материалы", recipient_bin=CP_BIN, **extra):
    imp = common.loaded_import(common.org_account(), [
        {"amount": amount, "purpose": purpose, "recipient_bin": recipient_bin, **extra}])
    return imp, imp.lines.get()


# ── кандидаты (ТЗ §11.2) ────────────────────────────────────────────────

def test_candidates_same_bin_close_amount_at_most_five_without_refused(slug):
    cp = common.counterparty(CP_BIN)
    close = [common.orm_invoice(cp, amount) for amount in (1000, 980, 1030, 950, 1090, 910)]
    far = common.orm_invoice(cp, 1200)
    refused = [common.orm_invoice(cp, 1000, status=status)
               for status in ("cancelled", "not_payable", "replaced")]
    usd = common.orm_invoice(cp, 1000, currency="USD")
    other = common.orm_invoice(common.counterparty("100000000099"), 1000)
    _, line = _line("1000")

    found = matching.candidates(line.pk)
    assert len(found) == 5
    assert [row["amount"] for row in found] == [D("1000.00"), D("980.00"), D("1030.00"),
                                                D("950.00"), D("1090.00")]
    assert all(row["same_bin"] for row in found)
    ids = {row["id"] for row in found}
    assert not ids & {str(inv.pk) for inv in (far, usd, other, *refused)}
    assert str(close[-1].pk) not in ids                      # шестой — за пределом пяти

    # Поиск: по сумме, по номеру (в «грязном» виде) и по контрагенту.
    assert [row["id"] for row in matching.candidates(line.pk, "1200")] == [str(far.pk)]
    dirty = "c" + far.number.lower()[1:]
    assert [row["id"] for row in matching.candidates(line.pk, dirty)] == [str(far.pk)]
    by_name = matching.candidates(line.pk, "Контрагент 099")
    assert [row["id"] for row in by_name] == [str(other.pk)]
    assert not by_name[0]["same_bin"]
    # Отказные статусы не находятся и поиском.
    assert matching.candidates(line.pk, refused[0].number) == []


# ── ручное сопоставление (ТЗ §11.4) ────────────────────────────────────

def test_manual_match_on_two_invoices_and_sum_over_line_is_422_on_amount(slug):
    cp = common.counterparty(CP_BIN)
    a, b = common.orm_invoice(cp, 600), common.orm_invoice(cp, 400)
    imp, line = _line("1000")

    err = _error(matching.match_line, FD, line.pk, [
        {"invoice_id": str(a.pk), "amount": "700"}, {"invoice_id": str(b.pk), "amount": "400"}])
    assert (err.code, err.fields[0]["field"]) == ("E-VAL-01", "amount")
    assert "больше суммы строки" in err.message

    line = matching.match_line(FD, line.pk, [
        {"invoice_id": str(a.pk), "amount": "600"}, {"invoice_id": str(b.pk), "amount": "400"}])
    assert line.match_status == LineMatchStatus.MATCHED
    assert {(m.invoice_id, m.amount, m.state, m.manual, m.created_by)
            for m in PaymentMatch.objects.filter(line=line)} == {
        (a.pk, D("600.00"), PaymentMatchState.ACTIVE, True, FD),
        (b.pk, D("400.00"), PaymentMatchState.ACTIVE, True, FD)}
    assert _recon(a) == (D("600.00"), ReconStatus.FULL)
    assert _recon(b) == (D("400.00"), ReconStatus.FULL)
    assert AuditLog.objects.filter(object_type="bpp.bankimport", object_id=str(imp.pk),
                                   action="line_matched", actor_id=FD).count() == 1
    # Уже сопоставленную строку вручную не сопоставить — сначала отмена.
    assert _error(matching.match_line, FD, line.pk,
                  [{"invoice_id": str(a.pk)}]).code == "E-STATE-01"


def test_manual_match_refuses_cancelled_not_payable_and_foreign_currency(slug):
    cp = common.counterparty(CP_BIN)
    _, line = _line("1000")
    for status in ("cancelled", "not_payable"):
        inv = common.orm_invoice(cp, 1000, status=status)
        err = _error(matching.match_line, FD, line.pk, [{"invoice_id": str(inv.pk)}])
        assert (err.code, err.fields[0]["field"]) == ("E-VAL-01", "invoice_id")
    usd = common.orm_invoice(cp, 1000, currency="USD")
    err = _error(matching.match_line, FD, line.pk, [{"invoice_id": str(usd.pk)}])
    assert (err.code, err.fields[0]["field"], err.message) == (
        "E-VAL-01", "invoice_id", "Валюта платежа ≠ валюте счёта")
    missing = _error(matching.match_line, FD, line.pk,
                     [{"invoice_id": "00000000-0000-0000-0000-000000000000"}])
    assert (missing.code, missing.fields[0]["field"]) == ("E-VAL-01", "invoice_id")
    line.refresh_from_db()
    assert line.match_status == LineMatchStatus.UNMATCHED
    assert not PaymentMatch.objects.exists()


# ── подтверждение и отмена (AC-011, BR-073, Review Focus 3) ────────────

def test_confirm_review_changes_status_cancel_returns_it_and_automatch_keeps_off(slug):
    cp = common.counterparty(CP_BIN)
    inv = common.orm_invoice(cp, 1000)
    imp, line = _line("400", purpose=f"Оплата {inv.number}", recipient_bin=FOREIGN_BIN)
    matching.auto_match(imp.pk)
    line.refresh_from_db()
    assert line.match_status == LineMatchStatus.NEEDS_REVIEW
    assert _recon(inv) == (D("0.00"), ReconStatus.NO_DATA)

    err = _error(matching.confirm, FD, line.pk, SHORT)
    assert (err.code, err.status, err.fields[0]["field"]) == ("BR-060", 422, "comment")

    line = matching.confirm(FD, line.pk, REASON)
    assert line.match_status == LineMatchStatus.MATCHED
    match = PaymentMatch.objects.get(line=line)
    assert (match.state, match.confirmed_by_id, match.comment, match.review_reason) == (
        PaymentMatchState.ACTIVE, FD, REASON, "bin_mismatch")
    assert match.confirmed_at is not None
    assert _recon(inv) == (D("400.00"), ReconStatus.PARTIAL)

    # «Отменить» — тоже с комментарием (BR-060).
    err = _error(matching.cancel_match, FD, line.pk, SHORT)
    assert (err.code, err.fields[0]["field"]) == ("BR-060", "comment")
    line = matching.cancel_match(FD, line.pk, comment="Платёж по другому счёту")
    assert line.match_status == LineMatchStatus.UNMATCHED
    match.refresh_from_db()
    assert (match.state, match.cancelled_by_id) == (PaymentMatchState.CANCELLED, FD)
    assert _recon(inv) == (D("0.00"), ReconStatus.NO_DATA)

    # Отменённое ФД сопоставление автосверка не возвращает.
    matching.auto_match(imp.pk)
    line.refresh_from_db()
    assert line.match_status == LineMatchStatus.UNMATCHED
    assert PaymentMatch.objects.filter(line=line).count() == 1
    assert _recon(inv) == (D("0.00"), ReconStatus.NO_DATA)
    # ... а вручную — можно.
    matching.match_line(FD, line.pk, [{"invoice_id": str(inv.pk)}])
    assert _recon(inv) == (D("400.00"), ReconStatus.PARTIAL)


def test_confirm_refuses_refused_invoice_status_and_foreign_currency(slug):
    cp = common.counterparty(CP_BIN)
    cancelled = common.orm_invoice(cp, 500, status="cancelled")
    usd = common.orm_invoice(cp, 700, currency="USD")
    imp = common.loaded_import(common.org_account(), [
        {"amount": "500", "purpose": f"Оплата {cancelled.number}", "recipient_bin": CP_BIN},
        {"amount": "700", "purpose": f"Оплата {usd.number}", "recipient_bin": CP_BIN},
    ])
    matching.auto_match(imp.pk)
    by_row = {line.row_no: line for line in imp.lines.all()}
    assert (by_row[1].review_reason, by_row[2].review_reason) == ("invoice_status",
                                                                  "currency_mismatch")

    err = _error(matching.confirm, FD, by_row[1].pk, REASON)
    assert (err.code, err.fields[0]["field"]) == ("E-VAL-01", "invoice_id")
    assert cancelled.number in err.message and "Отменён" in err.message
    err = _error(matching.confirm, FD, by_row[2].pk, REASON)
    assert (err.code, err.fields[0]["field"], err.message) == (
        "E-VAL-01", "invoice_id", "Валюта платежа ≠ валюте счёта")
    assert not PaymentMatch.objects.filter(state=PaymentMatchState.ACTIVE).exists()
    assert _recon(usd) == (D("0.00"), ReconStatus.NO_DATA)


# ── исключение (ТЗ §11.4) ───────────────────────────────────────────────

def test_exclude_needs_comment_and_leaves_the_other_tabs(slug):
    imp, line = _line("250", purpose="Комиссия банка")
    matching.auto_match(imp.pk)

    err = _error(matching.exclude, FD, line.pk, "")
    assert (err.code, err.status) == ("BR-060", 422)
    line = matching.exclude(FD, line.pk, "Комиссия банка за обслуживание")
    assert (line.match_status, line.excluded_by_id) == (LineMatchStatus.EXCLUDED, FD)

    totals = matching.totals(imp.pk)
    assert totals["excluded"] == {"count": 1, "amount": D("250.00")}
    assert all(totals[tab]["count"] == 0 for tab in ("matched", "needs_review", "unmatched"))
    assert imports.lines(imp, tab="excluded")["total"] == 1
    assert imports.lines(imp, tab="unmatched")["total"] == 0
    assert _error(imports.lines, imp, tab="bogus").code == "E-VAL-01"
    # Исключённую строку повторно не исключить и вручную не сопоставить.
    assert _error(matching.exclude, FD, line.pk, REASON).code == "E-STATE-01"
    # «Отменить» снимает исключение.
    line = matching.cancel_match(FD, line.pk, "Это всё-таки оплата поставщику")
    assert (line.match_status, line.excluded_at) == (LineMatchStatus.UNMATCHED, None)


def test_automatch_leaves_an_excluded_line_alone(slug):
    """M-3: строку, исключённую ФД, автосверка («Сверить») не трогает, даже
    если в назначении номер существующего счёта."""
    inv = common.orm_invoice(common.counterparty(CP_BIN), 1000)
    imp, line = _line("1000", purpose=f"Возврат аванса по {inv.number}")
    matching.exclude(FD, line.pk, "Возврат аванса, не закупка")

    result = matching.auto_match(imp.pk)

    line.refresh_from_db()
    assert (line.match_status, line.found_numbers) == (LineMatchStatus.EXCLUDED, [])
    assert not PaymentMatch.objects.exists()
    assert _recon(inv) == (D("0.00"), ReconStatus.NO_DATA)
    assert result["excluded"] == {"count": 1, "amount": D("1000.00")}


def test_exact_split_with_a_foreign_bin_invoice_sends_the_whole_line_to_review(slug):
    """M-5: несколько номеров, сумма делится ровно, но у одного счёта чужой
    БИН, у другого статус не «К оплате» — на проверку уходит вся строка. У
    каждого сопоставления своя причина, у счёта без замечаний — причина
    строки (первая по порядку номеров); статусы сверки не меняются, пока ФД
    не подтвердит, после подтверждения — все три счёта."""
    own = common.counterparty(CP_BIN)
    ok = common.orm_invoice(own, 300)
    foreign = common.orm_invoice(common.counterparty(FOREIGN_BIN), 500)
    pending = common.orm_invoice(own, 200, status="under_review")
    imp, line = _line("1000",
                      purpose=f"Оплата {ok.number}, {foreign.number} и {pending.number}")

    matching.auto_match(imp.pk)

    line.refresh_from_db()
    assert (line.match_status, line.review_reason) == (LineMatchStatus.NEEDS_REVIEW,
                                                       "bin_mismatch")
    assert line.found_numbers == [ok.number, foreign.number, pending.number]
    assert {(m.invoice_id, m.amount, m.state, m.review_reason) for m in line.matches.all()} == {
        (ok.pk, D("300.00"), PaymentMatchState.REVIEW, "bin_mismatch"),
        (foreign.pk, D("500.00"), PaymentMatchState.REVIEW, "bin_mismatch"),
        (pending.pk, D("200.00"), PaymentMatchState.REVIEW, "invoice_status")}
    for inv in (ok, foreign, pending):
        assert _recon(inv) == (D("0.00"), ReconStatus.NO_DATA)

    matching.confirm(FD, line.pk, REASON)

    assert _recon(ok) == (D("300.00"), ReconStatus.FULL)
    assert _recon(foreign) == (D("500.00"), ReconStatus.FULL)
    assert _recon(pending) == (D("200.00"), ReconStatus.FULL)


# ── отмена загрузки (Review Focus 5, I-1) ──────────────────────────────

def test_cancel_import_recalcs_frees_the_buh_mark_and_reupload_reconciles(
        slug, django_capture_on_commit_callbacks):
    _, _, inv = _payable(slug, D("1000"))
    buh = invoice_flow._buh(slug)
    inv = payments.mark_paid(buh, inv.id, pay_date=timezone.localdate(), amount=1000)
    account = common.org_account()
    docs = [_doc(account, "601", "1000.00", f"Оплата по счёту {inv.number}")]
    imp = _upload(django_capture_on_commit_callbacks, account, docs)
    assert imp.status == BankImportStatus.RECONCILED
    assert _recon(inv) == (D("1000.00"), ReconStatus.FULL)
    mark = inv.payments.get(cancelled_at__isnull=True)
    assert _error(payments.unmark, buh, inv.id, mark.pk,
                  comment="Ошибка в сумме платежа").status == 409

    assert matching.impact(imp.pk) == {"invoices": 1, "lines": 1}
    cancelled = matching.cancel_import(FD, imp.pk, comment="Выписка не того периода")
    assert cancelled.status == BankImportStatus.CANCELLED
    assert not imp.lines.filter(cancelled_at__isnull=True).exists()
    assert set(PaymentMatch.objects.values_list("state", flat=True)) == {
        PaymentMatchState.CANCELLED}
    assert _recon(inv) == (D("0.00"), ReconStatus.NO_DATA)
    # «Оплачено по банку» счёта = Σ единого определения подтверждённых сопоставлений.
    total = recon.active_matches().filter(invoice=inv).aggregate(t=Sum("amount"))["t"]
    assert (total or D("0")) == inv.paid_bank_amount
    # Отметка БУХ снова отменяема, повтор отмены загрузки — 409.
    payments.unmark(buh, inv.id, mark.pk, comment="Ошибка в сумме платежа")
    assert _error(matching.cancel_import, FD, imp.pk).code == "E-STATE-01"

    again = _upload(django_capture_on_commit_callbacks, account, docs)
    assert (again.status, again.duplicates) == (BankImportStatus.RECONCILED, 0)
    assert again.lines.get().match_status == LineMatchStatus.MATCHED
    assert _recon(inv) == (D("1000.00"), ReconStatus.FULL)
    assert AuditLog.objects.filter(object_type="bpp.bankimport", object_id=str(imp.pk),
                                   action="cancelled", actor_id=FD).count() == 1


def test_actions_on_a_cancelled_import_are_409(slug):
    imp, line = _line("100")
    matching.cancel_import(FD, imp.pk)
    for call in (lambda: matching.exclude(FD, line.pk, REASON),
                 lambda: matching.cancel_match(FD, line.pk, REASON),
                 lambda: matching.confirm(FD, line.pk, REASON),
                 lambda: matching.match_line(FD, line.pk, [{"invoice_id": str(line.pk)}]),
                 lambda: matching.reconcile(FD, imp.pk)):
        assert _error(call).code == "E-STATE-01"


# ── HTTP: права, идемпотентность, «Сверить», выгрузка ─────────────────

def _post(client, url, slug, user, body=None, key=None):
    headers = s.auth(slug, user)
    if key:
        headers["HTTP_IDEMPOTENCY_KEY"] = key
    return client.post(url, body or {}, **headers)


def test_rights_buh_reads_and_exports_but_cannot_act_td_cannot_read(slug):
    cp = common.counterparty(CP_BIN)
    inv = common.orm_invoice(cp, 1000)
    imp, line = _line("1000")
    client = Client()

    for user in (FD, BUH):
        headers = s.auth(slug, user)
        assert client.get(f"{BASE}/imports/{imp.pk}/lines?tab=unmatched",
                          **headers).status_code == 200
        candidates = client.get(f"{BASE}/lines/{line.pk}/candidates", **headers)
        assert candidates.status_code == 200
        assert [row["id"] for row in candidates.json()["items"]] == [str(inv.pk)]
        assert client.get(f"{BASE}/imports/{imp.pk}/impact", **headers).status_code == 200
        assert client.get(f"{BASE}/imports/{imp.pk}/export", **headers).status_code == 200

    for user in (BUH, TD):
        for url, body in ((f"{BASE}/lines/{line.pk}/exclude", {"comment": REASON}),
                          (f"{BASE}/lines/{line.pk}/match",
                           {"allocations": [{"invoice_id": str(inv.pk)}]}),
                          (f"{BASE}/lines/{line.pk}/cancel-match", {}),
                          (f"{BASE}/lines/{line.pk}/confirm", {"comment": REASON}),
                          (f"{BASE}/imports/{imp.pk}/cancel", {}),
                          (f"{BASE}/imports/{imp.pk}/reconcile", {})):
            assert _post(client, url, slug, user, body).status_code == 403, (user, url)
    headers = s.auth(slug, TD)
    assert client.get(f"{BASE}/imports/{imp.pk}/lines", **headers).status_code == 403
    assert client.get(f"{BASE}/lines/{line.pk}/candidates", **headers).status_code == 403
    assert client.get(f"{BASE}/imports/{imp.pk}/export", **headers).status_code == 403

    response = _post(client, f"{BASE}/lines/{line.pk}/match", slug, FD,
                     {"allocations": [{"invoice_id": str(inv.pk)}]})
    assert response.status_code == 200, response.json()
    body = response.json()
    assert body["match_status"] == "matched"
    assert body["matches"][0]["invoice_number"] == inv.number
    assert body["matches"][0]["paid_bank_amount"] == "1000.00"
    assert body["matches"][0]["recon_status"] == "full"


def test_repeated_post_with_the_same_key_is_one_action_and_one_journal_record(slug):
    imp, line = _line("300")
    client = Client()
    url = f"{BASE}/lines/{line.pk}/exclude"
    first = _post(client, url, slug, FD, {"comment": REASON}, key="exclude-line-1")
    second = _post(client, url, slug, FD, {"comment": REASON}, key="exclude-line-1")
    assert (first.status_code, second.status_code) == (200, 200)
    assert first.json() == second.json()
    assert first.json()["match_status"] == "excluded"
    with use_company(slug):
        assert AuditLog.objects.filter(object_type="bpp.bankimport", object_id=str(imp.pk),
                                       action="line_excluded").count() == 1
    # Без ключа — новое действие, и строка уже исключена.
    third = _post(client, url, slug, FD, {"comment": REASON})
    assert (third.status_code, third.json()["code"]) == (409, "E-STATE-01")
    short = _post(client, f"{BASE}/lines/{line.pk}/confirm", slug, FD, {"comment": SHORT})
    assert (short.status_code, short.json()["code"]) == (422, "BR-060")


def test_reconcile_button_runs_the_automatch_of_a_loaded_import(slug):
    """M-2: автосверка упала — загрузка осталась «Загружена»; «Сверить»
    доводит её. На «Сверена» — 409."""
    cp = common.counterparty(CP_BIN)
    inv = common.orm_invoice(cp, 1000)
    imp, _ = _line("1000", purpose=f"Оплата {inv.number}")
    assert imp.status == BankImportStatus.LOADED
    client = Client()
    url = f"{BASE}/imports/{imp.pk}/reconcile"

    response = _post(client, url, slug, FD)
    assert response.status_code == 200, response.json()
    card = response.json()
    assert (card["status"], card["totals"]["matched"]["count"]) == ("reconciled", 1)
    with use_company(slug):
        assert _recon(inv) == (D("1000.00"), ReconStatus.FULL)
        assert AuditLog.objects.filter(object_type="bpp.bankimport", object_id=str(imp.pk),
                                       action="reconciled", actor_id=FD).count() == 1
    again = _post(client, url, slug, FD)
    assert (again.status_code, again.json()["code"]) == (409, "E-STATE-01")


def test_export_has_three_sheets_and_amounts_are_numbers(slug):
    cp = common.counterparty(CP_BIN)
    inv = common.orm_invoice(cp, 1000)
    imp = common.loaded_import(common.org_account(), [
        {"amount": "600", "purpose": f"Оплата {inv.number}", "recipient_bin": CP_BIN},
        {"amount": "150.50", "purpose": f"Оплата {inv.number}", "recipient_bin": FOREIGN_BIN},
        {"amount": "75", "purpose": "Прочее", "recipient_bin": CP_BIN},
    ])
    matching.auto_match(imp.pk)

    response = Client().get(f"{BASE}/imports/{imp.pk}/export", **s.auth(slug, BUH))
    assert response.status_code == 200
    assert response["Content-Type"].startswith("application/vnd.openxmlformats")
    book = openpyxl.load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["Сопоставлены", "Требуют проверки", "Не сопоставлены"]

    def table(name):
        rows = list(book[name].iter_rows(values_only=True))
        return [dict(zip(rows[0], row, strict=True)) for row in rows[1:]]

    matched, review, unmatched = (table(name) for name in book.sheetnames)
    assert len(matched) == 1 and matched[0]["Счёт"] == inv.number
    assert matched[0]["Сумма"] == 600 and isinstance(matched[0]["Сумма"], (int, float))
    assert matched[0]["Оплачено по банку всего"] == 600
    assert review[0]["Сумма сопоставления"] == 150.5
    assert review[0]["Причина"] == "БИН получателя не совпадает с БИН контрагента счёта"
    assert [row["Сумма"] for row in unmatched] == [75]


def test_line_views_after_http_read_under_company(slug):
    """Строки по вкладкам через ручку — у сопоставления ссылка на счёт и
    «Оплачено по банку всего»."""
    cp = common.counterparty(CP_BIN)
    inv = common.orm_invoice(cp, 1000)
    imp, _ = _line("1000", purpose=f"Оплата {inv.number}")
    matching.auto_match(imp.pk)
    page = Client().get(f"{BASE}/imports/{imp.pk}/lines?tab=matched",
                        **s.auth(slug, FD)).json()
    assert page["total"] == 1
    match = page["items"][0]["matches"][0]
    assert (match["invoice_url"], match["invoice_amount"], match["paid_bank_amount"]) == (
        f"/bpp/invoices/{inv.pk}", "1000.00", "1000.00")
    with use_company(slug):
        assert BankStatementLine.objects.get(bank_import=imp).match_status == "matched"
        assert Invoice.objects.get(pk=inv.pk).recon_status == ReconStatus.FULL
