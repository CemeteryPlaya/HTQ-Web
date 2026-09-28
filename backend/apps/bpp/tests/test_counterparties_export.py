"""Выгрузка справочника «Контрагенты» в xlsx — контракт ``?format=xlsx``
поверх ``GET counterparties`` (ТЗ §19, D-32, задача 6).

Реестр и выгрузка делят один фильтр
(``services/counterparties/service.py::_filtered``/``_ordered``), поэтому
здесь проверяется сама ручка (три границы ``export.respond`` через реальные
строки контрагентов, а не синтетический генератор, как в
``test_export.py``) и то, что ``export_rows`` — функция пересборки фона —
не может разойтись со страницей экрана.
"""

from __future__ import annotations

import io

import openpyxl
import pytest
from django.test import Client

from apps.bpp.models import ExportJob
from apps.bpp.services.core import export
from apps.bpp.services.counterparties import service
from apps.bpp.tests import stage2 as s
from apps.bpp.tests.counterparties import common
from htqweb.tenancy.db import use_company

pytestmark = pytest.mark.django_db

BASE = "/api/bpp/v1/counterparties"
FD, NOBODY = 941, 942

HEADERS = ["Наименование", "БИН/ИИН", "Страна", "Статус", "Проверенный",
           "Удачных документов", "Дата создания"]


@pytest.fixture
def slug(company_context):
    common.countries()
    slug = company_context["slug"]
    s.grant(slug, FD, "bpp-fd")
    s.user(NOBODY)
    return slug


def _seed(slug, *, blocked_name: str | None = None):
    """Три контрагента: два в KZ (один из них может быть заблокирован
    ``blocked_name``) и один нерезидент в RU — тот же набор, что фильтрует
    ``counterparties/test_api.py::test_registry_search_filters_and_page_size``."""
    numbers = common.valid_bins(2)
    made = []
    for number, name in zip(numbers, ("ТОО Альфа", "ТОО Бета")):
        cp = service.create(common.data(number, name=name, short_name=""), actor_id=FD)
        if name == blocked_name:
            cp = service.block(cp.pk, reason="нет оригиналов документов",
                               expected_version=cp.version, actor_id=FD)
        made.append(cp)
    made.append(service.create(
        common.data("7707083893", name="ПАО Гамма", kind="nonresident", country_code="RU"),
        actor_id=FD))
    return made


def _load(response) -> openpyxl.Workbook:
    assert response["Content-Type"] == export.XLSX_MIME
    assert response["Content-Disposition"].startswith("attachment;")
    return openpyxl.load_workbook(io.BytesIO(response.content))


def test_small_selection_is_a_file_with_filters_applied(slug):
    _seed(slug, blocked_name="ТОО Бета")

    resp = Client().get(BASE, {"format": "xlsx", "country": "KZ"}, **s.auth(slug, FD))

    assert resp.status_code == 200
    ws = _load(resp).active
    assert [c.value for c in ws[1]] == HEADERS
    assert ws.max_row == 3  # заголовок + 2 контрагента KZ — RU в выборку не попал
    by_name = {ws.cell(row=r, column=1).value: r for r in (2, 3)}
    beta_row = by_name["ТОО Бета"]
    assert ws.cell(row=beta_row, column=4).value == "Заблокирован"
    assert ws.cell(row=beta_row, column=5).value == "Нет"
    alpha_row = by_name["ТОО Альфа"]
    assert ws.cell(row=alpha_row, column=3).value == "KZ"
    assert ws.cell(row=alpha_row, column=4).value == "Активен"
    assert ws.cell(row=alpha_row, column=6).value == "0"  # удачных документов пока нет


def test_search_filter_is_respected_in_the_export(slug):
    _seed(slug)

    resp = Client().get(BASE, {"format": "xlsx", "q": "Гамма"}, **s.auth(slug, FD))

    ws = _load(resp).active
    assert ws.max_row == 2  # заголовок + один нерезидент
    assert ws.cell(row=2, column=1).value == "ПАО Гамма"
    assert ws.cell(row=2, column=3).value == "RU"


def test_over_sync_limit_queues_a_job(slug, monkeypatch):
    """Review Focus 4 (по образцу ``test_export.py``): выборка больше
    ``SYNC_LIMIT`` — очередь, а не файл, и строка ``ExportJob`` в базе."""
    monkeypatch.setattr(export, "SYNC_LIMIT", 1)
    _seed(slug)  # 3 контрагента — больше подменённого предела

    resp = Client().get(BASE, {"format": "xlsx"}, **s.auth(slug, FD))

    assert resp.status_code == 200
    body = resp.json()
    assert body["status"] == "queued" and body["id"] and body["detail"]
    with use_company(slug):  # запрос вернул search_path в public
        job = ExportJob.objects.get(pk=body["id"])
    assert (job.row_count, job.requested_by) == (3, FD)


def test_over_hard_limit_is_422(slug, monkeypatch):
    monkeypatch.setattr(export, "HARD_LIMIT", 1)
    _seed(slug)  # 3 контрагента — больше подменённого потолка

    resp = Client().get(BASE, {"format": "xlsx"}, **s.auth(slug, FD))

    assert resp.status_code == 422
    assert resp.json()["code"] == "E-EXP-01"
    # Без use_company проверка смотрела бы в public и проходила бы всегда.
    with use_company(slug):
        assert not ExportJob.objects.exists()


def test_export_rows_matches_the_page_filter(slug):
    """``export_rows`` — функция пересборки фона: она обязана видеть ровно
    ту же выборку, что и ``registry()`` при тех же фильтрах (общий
    ``_filtered``). Колонки выгрузки деловые (БИН/ИИН), не ``id`` — сверяем
    по ``reg_number``, он уникален в паре (страна, номер)."""
    _seed(slug)
    filters = {"q": None, "countries": ["KZ"], "statuses": [], "sort": None}

    page_numbers = {row["reg_number"] for row in service.registry(**filters)["items"]}
    export_numbers = {row["reg_number"] for row in service.export_rows(**filters)}

    assert page_numbers and page_numbers == export_numbers


def test_export_requires_module_access(slug):
    resp = Client().get(BASE, {"format": "xlsx"}, **s.auth(slug, NOBODY))
    assert resp.status_code == 403
