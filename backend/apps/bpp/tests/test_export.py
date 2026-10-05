"""Экспорт реестров модуля в xlsx (``services/core/export.py``, задача 6).

Границы ТЗ §19 (Review Focus 4 плана этапа 2 A): ``SYNC_LIMIT`` (10 000)
строк — файл сразу; больше и до ``HARD_LIMIT`` (50 000) — фоном со ссылкой
в уведомлении; больше — 422 ``E-EXP-01``; пустая выборка — только
заголовки. Строки для граничных тестов — генератор без обращения к БД,
чтобы прогон был быстрым (10 000+ строк реестра не пишутся в тестовую базу).
Фоновая ветка идёт через ``transaction.on_commit`` —
``django_capture_on_commit_callbacks(execute=True)`` исполняет постановку
задачи, а ``CELERY_TASK_ALWAYS_EAGER`` (settings/test.py) — саму задачу.
"""

from __future__ import annotations

import io
from datetime import date, datetime, timezone

import openpyxl
import pytest
from django.test import Client

from apps.bpp.models import AuditLog, ExportJob, ExportStatus
from apps.bpp.services.core import export
from apps.bpp.tests import stage2 as s
from apps.bpp.tests.test_files import memory_storage  # noqa: F401  (фикстура)
from htqweb.errors import DomainError
from htqweb.tenancy.db import use_company

pytestmark = pytest.mark.django_db

BASE = "/api/bpp/v1"

COLUMNS = (
    export.Column("number", "Номер"),
    export.Column("amount", "Сумма", kind="money"),
    export.Column("created", "Создано", kind="date"),
)

PROBE = "apps.bpp.tests.test_export._rebuild_probe"


def _rows(n: int):
    for i in range(n):
        yield {"number": f"З-{i:05d}", "amount": f"{1000 + i}.50", "created": "2026-09-27"}


def _rebuild_probe(n: int):
    """Функция пересборки для фонового теста — путь до неё передаётся в
    ``rebuild`` строкой, как в реальном вызове из реестра."""
    return _rows(n)


def _rebuild_broken():
    raise RuntimeError("источник недоступен")


def _load(response) -> openpyxl.Workbook:
    assert response["Content-Type"] == export.XLSX_MIME
    assert response["Content-Disposition"].startswith("attachment;")
    return openpyxl.load_workbook(io.BytesIO(response.content))


def _background(request, count, rebuild, capture):
    with capture(execute=True):
        return export.respond(request, name="Большой реестр", columns=COLUMNS, rows=None,
                              count=count, rebuild=rebuild)


# ── файл сразу ──────────────────────────────────────────────────────────

def test_column_rejects_unknown_kind():
    with pytest.raises(ValueError):
        export.Column("x", "X", kind="nope")


def test_small_selection_is_returned_immediately(company_context):
    actor = s.actor(company_context["slug"], 1, "bpp-sn")

    response = export.respond(actor.request, name="Реестр", columns=COLUMNS,
                              rows=_rows(3), count=3)

    ws = _load(response).active
    assert [c.value for c in ws[1]] == ["Номер", "Сумма", "Создано"]
    assert ws.max_row == 4  # заголовок + 3 строки
    row2 = [c.value for c in ws[2]]
    assert row2[0] == "З-00000"
    assert row2[1] == 1000.5  # money — число, не строка
    assert ws.cell(row=2, column=2).number_format == "#,##0.00"
    # date — датой (openpyxl читает её как datetime на полночь), не строкой
    assert row2[2].date() == date(2026, 9, 27)
    assert ws.cell(row=2, column=3).number_format == "dd.mm.yyyy"
    assert not ExportJob.objects.exists()


def test_empty_selection_is_headers_only(company_context):
    actor = s.actor(company_context["slug"], 1, "bpp-sn")

    response = export.respond(actor.request, name="Реестр", columns=COLUMNS,
                              rows=(), count=0)

    ws = _load(response).active
    assert ws.max_row == 1
    assert [c.value for c in ws[1]] == ["Номер", "Сумма", "Создано"]


def test_exactly_sync_limit_is_returned_immediately(company_context):
    """Review Focus 4: 10 000 строк — файл сразу."""
    actor = s.actor(company_context["slug"], 1, "bpp-sn")

    response = export.respond(actor.request, name="Реестр", columns=COLUMNS,
                              rows=_rows(export.SYNC_LIMIT), count=export.SYNC_LIMIT)

    assert _load(response).active.max_row == export.SYNC_LIMIT + 1
    assert not ExportJob.objects.exists()


def test_cyrillic_name_goes_into_filename_star():
    """Имя реестра — кириллица, заголовок HTTP — latin-1: ``filename*``."""
    response = export._xlsx_response("Реестр заявок", COLUMNS, ())
    assert "filename*=utf-8''" in response["Content-Disposition"]


def test_datetime_is_written_in_platform_time_zone():
    """Excel не хранит пояс: время пишется так, как его читает человек."""
    columns = (export.Column("at", "Когда", kind="datetime"),)
    rows = [{"at": datetime(2026, 9, 27, 20, 30, tzinfo=timezone.utc)},
            {"at": "2026-09-27T20:30:00Z"}]

    ws = openpyxl.load_workbook(io.BytesIO(export.write_xlsx("Р", columns, rows))).active

    assert ws.cell(row=2, column=1).value == datetime(2026, 9, 28, 1, 30)
    assert ws.cell(row=3, column=1).value == datetime(2026, 9, 28, 1, 30)
    assert ws.cell(row=2, column=1).number_format == "dd.mm.yyyy hh:mm"


def test_integer_is_a_whole_number_without_decimals():
    """Счётчики (строк, дублей) — целым числом с разрядами, без «,00»;
    не число — текстом, а не падением выгрузки."""
    columns = (export.Column("n", "Строк", kind="integer"),)
    rows = [{"n": 1250}, {"n": "17"}, {"n": None}, {"n": "много"}]

    ws = openpyxl.load_workbook(io.BytesIO(export.write_xlsx("Р", columns, rows))).active

    assert [ws.cell(row=r, column=1).value for r in range(2, 6)] == [1250, 17, None, "много"]
    assert ws.cell(row=2, column=1).number_format == "#,##0"
    assert isinstance(ws.cell(row=2, column=1).value, int)


def test_text_starting_with_equals_is_not_a_formula():
    """Наименование «=HYPERLINK(…)» не должно выполниться у читателя файла."""
    columns = (export.Column("name", "Наименование"),)

    ws = openpyxl.load_workbook(
        io.BytesIO(export.write_xlsx("Р", columns, [{"name": "=1+2"}]))).active

    cell = ws.cell(row=2, column=1)
    assert (cell.value, cell.data_type) == ("=1+2", "s")


def test_sheet_title_drops_characters_excel_forbids():
    data = export.write_xlsx("Реестр: счета [2026]/сентябрь", COLUMNS, ())
    assert openpyxl.load_workbook(io.BytesIO(data)).active.title == "Реестр счета 2026сентябрь"


def test_write_limit_is_enforced_while_streaming():
    """Фон пересобирает выборку заново — строк могло стать больше предела."""
    with pytest.raises(DomainError) as exc:
        export.write_xlsx("Р", COLUMNS, _rows(3), limit=2)
    assert exc.value.code == "E-EXP-01"


def test_write_xlsx_sheets_keeps_column_formats_and_unique_titles():
    """Книга из нескольких листов («Экспорт результата» сверки выписки):
    форматы колонок те же, что у одного листа; имена листов — по правилам
    Excel и без повторов; без листов — книга с одним пустым листом."""
    rows = [{"number": "СЧ-1", "amount": "1250.50", "created": "2026-09-27"}]
    data = export.write_xlsx_sheets("Книга", [("Сопоставлены", COLUMNS, rows),
                                              ("Требуют: проверки", COLUMNS, ()),
                                              ("Сопоставлены", COLUMNS, ())])

    book = openpyxl.load_workbook(io.BytesIO(data))
    assert book.sheetnames == ["Сопоставлены", "Требуют проверки", "Сопоставлены 2"]
    ws = book["Сопоставлены"]
    assert [cell.value for cell in ws[1]] == ["Номер", "Сумма", "Создано"]
    assert (ws.cell(row=2, column=2).value, ws.cell(row=2, column=2).number_format) == (
        1250.5, "#,##0.00")
    assert ws.cell(row=2, column=3).value == datetime(2026, 9, 27)
    assert book["Требуют проверки"].max_row == 1
    empty = openpyxl.load_workbook(io.BytesIO(export.write_xlsx_sheets("Пусто", [])))
    assert empty.sheetnames == ["Пусто"]


# ── отказ и ошибки вызывающего ──────────────────────────────────────────

def test_over_hard_limit_is_422(company_context):
    """Review Focus 4: 50 001 (> HARD_LIMIT) — отказ до всякой работы."""
    actor = s.actor(company_context["slug"], 1, "bpp-sn")

    with pytest.raises(DomainError) as exc:
        export.respond(actor.request, name="Реестр", columns=COLUMNS, rows=None,
                       count=export.HARD_LIMIT + 1, rebuild=(PROBE, {"n": 1}))

    assert (exc.value.code, exc.value.status) == ("E-EXP-01", 422)
    assert str(export.HARD_LIMIT + 1) in exc.value.message
    assert not ExportJob.objects.exists()


@pytest.mark.parametrize("rebuild", [
    None,                                          # фон без функции пересборки
    ("apps.users.services.export_rows", {}),       # функция вне модуля
    (PROBE, {"since": date(2026, 9, 27)}),         # фильтр не везётся брокером
])
def test_background_misuse_is_a_programming_error(company_context, rebuild):
    actor = s.actor(company_context["slug"], 1, "bpp-sn")
    with pytest.raises(ValueError):
        export.respond(actor.request, name="Реестр", columns=COLUMNS, rows=None,
                       count=export.SYNC_LIMIT + 1, rebuild=rebuild)
    assert not ExportJob.objects.exists()


# ── фон ─────────────────────────────────────────────────────────────────

def test_background_job_stores_file_and_notifies(company_context,
                                                 django_capture_on_commit_callbacks):
    """Review Focus 4: 10 001 строка — фоном, ссылка в уведомлении."""
    from apps.media_files.models import FileMetadata
    from apps.notifications.models import Notification

    actor = s.actor(company_context["slug"], 42, "bpp-sn")
    count = export.SYNC_LIMIT + 1

    result = _background(actor.request, count, (PROBE, {"n": count}),
                         django_capture_on_commit_callbacks)

    assert result["status"] == ExportStatus.QUEUED
    job = ExportJob.objects.get(pk=result["id"])
    assert (job.status, job.requested_by, job.row_count) == (ExportStatus.DONE, 42, count)
    assert job.finished_at is not None

    meta = FileMetadata.objects.get(pk=job.media_file_id)
    assert (meta.scope, meta.owner_id, meta.is_public) == ("generic", 42, False)
    assert meta.mime == export.XLSX_MIME

    note = Notification.objects.get(target_type="bpp.export", target_id=str(job.pk))
    assert note.recipient_id == 42
    assert note.company_slug == company_context["slug"]
    assert note.url == f"/bpp/exports/{job.pk}"


def test_background_job_records_the_error_and_tells_the_requester(
        company_context, django_capture_on_commit_callbacks):
    """Пересборка упала — выгрузка получает читаемую причину, а заказчик —
    уведомление, а не вечное «готовится» (Celery eager в тесте перевыбросил
    бы исключение наружу, если бы задача его не поймала)."""
    from apps.notifications.models import Notification

    actor = s.actor(company_context["slug"], 43, "bpp-sn")

    result = _background(actor.request, export.SYNC_LIMIT + 1,
                         ("apps.bpp.tests.test_export._rebuild_broken", {}),
                         django_capture_on_commit_callbacks)

    job = ExportJob.objects.get(pk=result["id"])
    assert job.status == ExportStatus.ERROR
    assert job.error and job.media_file_id is None
    note = Notification.objects.get(target_type="bpp.export", target_id=str(job.pk))
    assert "не собран" in note.title


def test_background_job_is_not_rebuilt_twice(company_context,
                                             django_capture_on_commit_callbacks):
    """Повторная доставка задачи Celery не пересобирает готовую выгрузку."""
    from apps.bpp.tasks_export import build_export

    actor = s.actor(company_context["slug"], 44, "bpp-sn")
    count = export.SYNC_LIMIT + 1
    result = _background(actor.request, count, (PROBE, {"n": count}),
                         django_capture_on_commit_callbacks)
    first = ExportJob.objects.get(pk=result["id"]).media_file_id

    build_export.delay(company_slug=company_context["slug"], job_id=result["id"],
                       columns=[c.as_dict() for c in COLUMNS], rebuild_path=PROBE,
                       rebuild_kwargs={"n": count})

    assert ExportJob.objects.get(pk=result["id"]).media_file_id == first


# ── ручка GET exports/<id> ──────────────────────────────────────────────

def _job(slug, user_id, *, status=ExportStatus.QUEUED, media_file_id=None) -> ExportJob:
    return ExportJob.objects.create(requested_by=user_id, name="Реестр", row_count=10_001,
                                    status=status, media_file_id=media_file_id)


def test_download_gives_the_requester_a_link_and_journals_it(
        company_context, django_capture_on_commit_callbacks):
    slug = company_context["slug"]
    actor = s.actor(slug, s.SN, "bpp-sn")
    count = export.SYNC_LIMIT + 1
    result = _background(actor.request, count, (PROBE, {"n": count}),
                         django_capture_on_commit_callbacks)

    resp = Client().get(f"{BASE}/exports/{result['id']}", HTTP_USER_AGENT="pytest-agent",
                        **s.auth(slug, s.SN))

    assert resp.status_code == 200
    body = resp.json()
    assert body["status"] == "done"
    assert body["url"] and body["expires_at"]
    with use_company(slug):  # запрос вернул search_path в public
        entry = AuditLog.objects.get(object_type=export.AUDIT_OBJECT_TYPE,
                                     object_id=result["id"])
    assert (entry.action, entry.actor_id) == ("file_downloaded", s.SN)
    assert entry.changes["user_agent"] == "pytest-agent"
    assert entry.changes["ip"]


def test_download_of_a_job_in_progress_has_no_link(company_context):
    slug = company_context["slug"]
    s.grant(slug, s.SN, "bpp-sn")
    job = _job(slug, s.SN)

    resp = Client().get(f"{BASE}/exports/{job.pk}", **s.auth(slug, s.SN))

    assert resp.status_code == 200
    assert resp.json()["status"] == "queued"
    assert "url" not in resp.json()
    assert not AuditLog.objects.filter(object_id=str(job.pk)).exists()


def test_someone_elses_job_is_404(company_context):
    slug = company_context["slug"]
    s.grant(slug, s.SN, "bpp-sn")
    s.grant(slug, s.PM, "bpp-pm")
    job = _job(slug, s.PM)

    resp = Client().get(f"{BASE}/exports/{job.pk}", **s.auth(slug, s.SN))

    assert resp.status_code == 404


@pytest.mark.parametrize("job_id", ["not-a-uuid", "00000000-0000-4000-8000-000000000000"])
def test_bad_or_unknown_id_is_404(company_context, job_id):
    slug = company_context["slug"]
    s.grant(slug, s.SN, "bpp-sn")

    resp = Client().get(f"{BASE}/exports/{job_id}", **s.auth(slug, s.SN))

    assert resp.status_code == 404
    assert "detail" in resp.json()


def test_download_requires_module_access(company_context):
    slug = company_context["slug"]
    job = _job(slug, 45)

    resp = Client().get(f"{BASE}/exports/{job.pk}", **s.auth(slug, 45))

    assert resp.status_code == 403
