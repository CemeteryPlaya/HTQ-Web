"""Отчёт R-01 «KPI снабжения» и записи KPI (план этапа 5 A, задача 5).

Записи KPI и АП пишутся прямо в базу (``kpi_helpers``): отчёту важны
статусы, суммы, даты и авторы. Права и охват СН проверяются по HTTP —
через настоящий гейт.
"""

from __future__ import annotations

import io
from datetime import timedelta
from zoneinfo import ZoneInfo

import openpyxl
import pytest
from django.conf import settings
from django.test import Client
from django.utils import timezone

from apps.bpp.models import KpiStatus, OfferStatus
from apps.bpp.services.alternatives import report
from apps.bpp.tests import stage2 as s
from apps.bpp.tests.alternatives.kpi_helpers import SN_A, SN_B, invoice, kpi_row, submitted_offer, users

pytestmark = pytest.mark.django_db
BASE = "/api/bpp/v1/kpi"
GD, BUH = 910, 907


def _fd(slug):
    return s.actor(slug, s.FD, "bpp-fd")


def _rows(data) -> dict:
    return {row["buyer_id"]: row for row in data["rows"]}


def _get(slug, user_id, code, url, **params):
    s.grant(slug, user_id, code)
    return Client().get(f"{BASE}/{url}", data=params, **s.auth(slug, user_id))


def _local_date(moment):
    return moment.astimezone(ZoneInfo(settings.PLATFORM_TIME_ZONE)).date()


# ── счётчики и доля ─────────────────────────────────────────────────────

def test_r01_counts_and_share(company_context):
    slug = company_context["slug"]
    users()
    src = invoice(1000)
    kpi_row(buyer_id=SN_A, saving="100.00")  # АП «Выбрано» + KPI «Подтверждён»
    submitted_offer(src, author_id=SN_A)  # «Подано», без KPI
    submitted_offer(src, author_id=SN_A, status=OfferStatus.WITHDRAWN)
    submitted_offer(src, author_id=SN_A, status=OfferStatus.ANNULLED)
    kpi_row(buyer_id=SN_B, saving="50.00")
    data = report.report(_fd(slug), report.Filters())
    rows = _rows(data)
    a, b = rows[SN_A], rows[SN_B]
    # «Отозвано» и «Аннулировано» в знаменатель KPI-004 не входят.
    assert (a["submitted"], a["selected"], a["confirmed"], a["share_pct"]) == (2, 1, 1, "50.00")
    assert (b["submitted"], b["selected"], b["confirmed"], b["share_pct"]) == (1, 1, 1, "100.00")
    total = data["total"]
    assert (total["submitted"], total["selected"], total["confirmed"]) == (3, 2, 2)
    assert total["share_pct"] == "66.67" and total["saving"] == "150.00"
    assert a["role"] == "sn" and a["name"]


def test_r01_saving_and_overspend_separate(company_context):
    slug = company_context["slug"]
    users()
    kpi_row(saving="100.00")
    kpi_row(saving="-30.00")  # удорожание: KPI-003, из KPI-002 не вычитается
    kpi_row(saving="999.00", status=KpiStatus.PRELIMINARY)
    kpi_row(saving="5.00", status=KpiStatus.ANNULLED)
    row = _rows(report.report(_fd(slug), report.Filters()))[SN_A]
    assert (row["saving"], row["overspend"], row["confirmed"], row["selected"]) == (
        "100.00", "30.00", 2, 4)


def test_r01_own_document_count(company_context):
    slug = company_context["slug"]
    users()
    kpi_row(own=True)
    kpi_row()
    assert _rows(report.report(_fd(slug), report.Filters()))[SN_A]["own_document_count"] == 1


def test_period_filters_by_selected_at(company_context):
    slug = company_context["slug"]
    users()
    kpi_row(buyer_id=SN_A, days_ago=1)
    kpi_row(buyer_id=SN_B, days_ago=10)
    day = _local_date(timezone.now() - timedelta(days=1))
    data = report.report(_fd(slug), report.Filters(period_from=day, period_to=day))
    a = _rows(data)[SN_A]
    # Выбрано вчера, а подано сегодня: за вчерашний период «подано» нет, доля — null.
    assert (a["selected"], a["submitted"], a["share_pct"]) == (1, 0, None)
    assert SN_B not in _rows(data)
    wide = report.report(_fd(slug), report.Filters(period_from=day - timedelta(days=30),
                                                   period_to=day))
    assert wide["total"]["selected"] == 2


def test_project_and_buyer_filters(company_context):
    slug = company_context["slug"]
    users()
    first = kpi_row(buyer_id=SN_A)
    kpi_row(buyer_id=SN_B)
    by_project = report.report(_fd(slug), report.Filters(project_id=str(first.project_id)))
    assert list(_rows(by_project)) == [SN_A]
    by_buyer = report.report(_fd(slug), report.Filters(buyer_id=SN_B))
    assert list(_rows(by_buyer)) == [SN_B]


# ── СН видит только своё ────────────────────────────────────────────────

def test_buyer_sees_only_own_rows_and_records(company_context):
    slug = company_context["slug"]
    users()
    mine = kpi_row(buyer_id=SN_A, saving="100.00")
    other = kpi_row(buyer_id=SN_B, saving="200.00")
    got = _get(slug, SN_A, "bpp-sn", "report")
    assert got.status_code == 200
    body = got.json()
    assert [row["buyer_id"] for row in body["rows"]] == [SN_A]
    assert body["total"]["saving"] == "100.00"
    # Чужой buyer_id в фильтре охват не расширяет.
    assert _get(slug, SN_A, "bpp-sn", "report", buyer_id=SN_B).json()["rows"] == []

    records = _get(slug, SN_A, "bpp-sn", "records").json()
    assert [item["id"] for item in records["items"]] == [str(mine.pk)] and records["total"] == 1
    assert _get(slug, SN_A, "bpp-sn", f"records/{other.pk}").status_code == 404
    assert _get(slug, SN_A, "bpp-sn", f"records/{mine.pk}").status_code == 200

    book = _get(slug, SN_A, "bpp-sn", "report/export")
    assert book.status_code == 200
    sheet = openpyxl.load_workbook(io.BytesIO(book.content))["Записи KPI"]
    assert sheet.max_row == 2  # заголовок и одна своя запись


def test_export_two_sheets_numeric(company_context):
    slug = company_context["slug"]
    users()
    kpi_row(buyer_id=SN_A, saving="100.00")
    kpi_row(buyer_id=SN_B, saving="200.50")
    response = _get(slug, s.FD, "bpp-fd", "report/export")
    assert response.status_code == 200
    book = openpyxl.load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["KPI снабжения", "Записи KPI"]
    summary, records = book["KPI снабжения"], book["Записи KPI"]
    assert summary.max_row == 4  # заголовок, два покупателя, итог
    savings = [summary.cell(row, 7).value for row in (2, 3, 4)]
    assert all(isinstance(value, (int, float)) for value in savings)
    assert sorted(savings) == [100, 200.5, 300.5]
    assert records.max_row == 3
    header = [cell.value for cell in records[1]]
    money = records.cell(2, header.index("Исходная часть, KZT") + 1).value
    assert isinstance(money, (int, float))


# ── права ───────────────────────────────────────────────────────────────

@pytest.mark.parametrize("user_id,code,status", [
    (s.TD, "bpp-td", 403), (BUH, "bpp-buh", 403), (s.OD, "bpp-od", 200),
    (GD, "bpp-gd", 200), (s.FD, "bpp-fd", 200)])
def test_report_access_by_role(company_context, user_id, code, status):
    slug = company_context["slug"]
    users()
    s.user(GD)
    kpi_row(buyer_id=SN_A)
    for url in ("report", "records", "report/export"):
        assert _get(slug, user_id, code, url).status_code == status, url


def test_od_and_gd_see_all_rows(company_context):
    slug = company_context["slug"]
    users()
    s.user(GD)
    kpi_row(buyer_id=SN_A)
    kpi_row(buyer_id=SN_B)
    for user_id, code in ((s.OD, "bpp-od"), (GD, "bpp-gd")):
        assert len(_get(slug, user_id, code, "report").json()["rows"]) == 2


def test_annul_endpoint_rights_and_result(company_context):
    slug = company_context["slug"]
    users()
    record = kpi_row(buyer_id=SN_A, status=KpiStatus.PRELIMINARY)
    url = f"{BASE}/records/{record.pk}/annul"
    body = {"comment": "Ошибочный выбор поставщика", "version": record.version}
    s.grant(slug, SN_A, "bpp-sn")
    assert Client().post(url, data=body, **s.auth(slug, SN_A)).status_code == 403
    s.grant(slug, s.FD, "bpp-fd")
    short = Client().post(url, data={"comment": "коротко"}, **s.auth(slug, s.FD))
    assert short.status_code == 422 and short.json()["code"] == "BR-060"
    ok = Client().post(url, data=body, **s.auth(slug, s.FD))
    assert ok.status_code == 200, ok.content
    assert ok.json()["status"] == "annulled" and ok.json()["annulled_by_id"] == s.FD


def test_records_filters_and_validation(company_context):
    slug = company_context["slug"]
    users()
    kpi_row(status=KpiStatus.CONFIRMED)
    kpi_row(status=KpiStatus.PRELIMINARY, own=True)
    got = _get(slug, s.FD, "bpp-fd", "records", status="preliminary").json()
    assert got["total"] == 1 and got["items"][0]["own_document"] is True
    assert got["items"][0]["source_url"].startswith("/bpp/invoices/")
    assert _get(slug, s.FD, "bpp-fd", "records", own_document="1").json()["total"] == 1
    assert _get(slug, s.FD, "bpp-fd", "records", status="zzz").status_code == 422
    assert _get(slug, s.FD, "bpp-fd", "report", period_from="вчера").status_code == 422
    assert _get(slug, s.FD, "bpp-fd", "records/not-a-uuid").status_code == 404
