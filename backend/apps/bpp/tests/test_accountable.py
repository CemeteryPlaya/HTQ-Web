"""Подотчётные средства (задача B4.1): резерв в «Задействовано», выдача,
авансовые отчёты — поведение как в contracts, источник — статья проекта."""

from __future__ import annotations

from decimal import Decimal

import io

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client
from openpyxl import load_workbook

from apps.bpp.models import AccountableStatus
from apps.bpp.services.accountable import accountable as service
from apps.bpp.services.accountable import read
from apps.bpp.services.budget import balance
from apps.bpp.services.budget import committed as calc
from apps.bpp.tests import stage2 as s
from apps.bpp.tests.test_files import PDF, memory_storage  # noqa: F401  (фикстура)
from apps.signoff import interface as signoff
from htqweb.errors import DomainError

pytestmark = pytest.mark.django_db
BUH = 907
BASE = "/api/bpp/v1"


def _routes():
    for subject in ("bpp.accountable_funds_request", "bpp.advance_report"):
        s.user(s.FD)
        signoff.configure_route(subject_type=subject, name="ФД", stages=[
            {"order": 1, "name": "ФД", "quorum": "any", "approver_kind": "users",
             "user_ids": [s.FD]}])


def _approve(subject: str, object_id) -> None:
    process = signoff.get_process_for(subject, str(object_id))
    task = next(t for st in process["stages"] for t in st["tasks"] if t["state"] == "pending")
    assert signoff.decide_many(actor_id=task["user_id"], items=[
        {"task_id": task["id"], "decision": "approve"}])[0]["ok"]


def _setup(slug, limit=100_000):
    proj, art = s.project(), s.metal()
    s.approved_budget(slug, proj, {art: limit})
    _routes()
    return proj, art, s.actor(slug, s.SN, "bpp-sn")


def _pdf():
    return SimpleUploadedFile("chek.pdf", PDF, content_type="application/pdf")


def test_amount_over_the_balance_is_e_bud_01(company_context):
    slug = company_context["slug"]
    proj, art, sn = _setup(slug, limit=1000)
    with pytest.raises(DomainError) as exc:
        service.create(sn, project_id=proj.id, article_id=art.id, amount=1500, goal="Командировка")
    assert exc.value.code == "E-BUD-01"


def test_reserve_starts_on_submit_and_stays_after_closing(company_context):
    slug = company_context["slug"]
    proj, art, sn = _setup(slug)
    req = service.create(sn, project_id=proj.id, article_id=art.id, amount=600,
                         goal="Командировка на объект")
    assert calc.committed_for(proj.id, art.id) == Decimal("0.00")
    service.submit(sn, req.id, expected_version=None)
    assert calc.committed_for(proj.id, art.id) == Decimal("600.00")
    assert calc.committed_reference(proj.id, art.id) == Decimal("600.00")
    assert balance.balance(proj.id, art.id)["available"] == Decimal("99400.00")


def test_full_path_to_closed_by_reports(company_context):
    slug = company_context["slug"]
    proj, art, sn = _setup(slug)
    req = service.create(sn, project_id=proj.id, article_id=art.id, amount=1000,
                         goal="Расходные материалы")
    req = service.submit(sn, req.id, expected_version=None)
    assert req.status == AccountableStatus.ON_REVIEW
    _approve("bpp.accountable_funds_request", req.id)
    req.refresh_from_db()
    assert req.status == AccountableStatus.AWAITING_ACCOUNTING

    with pytest.raises(DomainError) as exc:
        service.mark_paid(sn, req.id, expected_version=None)
    assert exc.value.status == 403
    buh = s.actor(slug, BUH, "bpp-buh")
    req = service.mark_paid(buh, req.id, expected_version=None)
    assert req.status == AccountableStatus.AWAITING_REPORT

    first = service.add_report(sn, req.id, expense_name="Бензин", amount=400, upload=_pdf())
    assert [r["can_submit"] for r in service.card(sn, req)["reports"]] == [True]
    service.submit_report(sn, first.id)
    # На согласовании отчёт второй раз не отправить; чужому — не отправить вовсе.
    assert [r["can_submit"] for r in service.card(sn, req)["reports"]] == [False]
    with pytest.raises(DomainError) as exc:
        service.add_report(sn, req.id, expense_name="Лишнее", amount=700, upload=_pdf())
    assert exc.value.code == "E-ACN-01"

    _approve("bpp.advance_report", first.id)
    req.refresh_from_db()
    assert req.status == AccountableStatus.AWAITING_REPORT
    second = service.add_report(sn, req.id, expense_name="Крепёж", amount=600, upload=_pdf())
    service.submit_report(sn, second.id)
    _approve("bpp.advance_report", second.id)
    req.refresh_from_db()
    assert req.status == AccountableStatus.CLOSED
    assert service.reported_amount(req) == Decimal("1000.00")
    # Закрытый подотчёт остаётся в «Задействовано»: выданные деньги потрачены.
    assert calc.committed_for(proj.id, art.id) == Decimal("1000.00")


def test_only_the_owner_and_fd_see_the_request(company_context):
    slug = company_context["slug"]
    proj, art, sn = _setup(slug)
    req = service.create(sn, project_id=proj.id, article_id=art.id, amount=10, goal="Мелочи")
    with pytest.raises(DomainError):
        service.get_visible(s.actor(slug, s.SN2, "bpp-sn"), req.id)
    assert service.get_visible(s.actor(slug, s.FD, "bpp-fd"), req.id) == req


def test_registry_is_paged_filtered_and_exported(company_context):
    """Реестр подотчёта — конверт реестров модуля (ТЗ §19): свои заявки,
    поиск по номеру и цели, итог по всей выборке; выгрузка — та же выборка."""
    slug = company_context["slug"]
    proj, art, sn = _setup(slug)
    mine = service.create(sn, project_id=proj.id, article_id=art.id, amount=100,
                          goal="Канцелярия для прорабской")
    service.create(sn, project_id=proj.id, article_id=art.id, amount=50, goal="Такси")

    page = read.registry(sn, filters={"search": "канц"})
    assert [row["number"] for row in page["items"]] == [mine.number]
    assert page["total"] == 1 and page["totals"]["amount"] == Decimal("100.00")
    row = page["items"][0]
    assert (row["project_code"], row["article_name"]) == ("П-015", art.name)
    assert row["reported_amount"] == Decimal("0.00") and row["remaining_amount"] == Decimal("100.00")
    assert read.registry(sn, filters={})["totals"]["amount"] == Decimal("150.00")
    assert read.registry(s.actor(slug, s.SN2, "bpp-sn"), filters={})["total"] == 0

    card = service.card(sn, mine)
    assert card["project"]["code"] == "П-015" and card["accountable_user_name"]

    client = Client()
    resp = client.get(f"{BASE}/accountable?search=Канц&page_size=25", **s.auth(slug, s.SN))
    assert resp.status_code == 200, resp.content
    assert resp.json()["page_size"] == 25 and resp.json()["total"] == 1
    sheet = client.get(f"{BASE}/accountable?search=Канц&format=xlsx", **s.auth(slug, s.SN))
    assert sheet.status_code == 200, sheet.content
    rows = list(load_workbook(io.BytesIO(sheet.content)).active.iter_rows(values_only=True))
    assert rows[0][0] == "Номер" and [r[0] for r in rows[1:]] == [mine.number]
