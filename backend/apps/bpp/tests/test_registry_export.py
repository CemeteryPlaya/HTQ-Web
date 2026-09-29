"""Выгрузка реестров бюджета, заявок и плана в xlsx и печать заявки
(ТЗ §19, D-32; остаток B этапа 2, B-3), узлы ``bpp.requests.all`` и
``bpp.plan.reassign`` (B-4)."""

from __future__ import annotations

import io
from decimal import Decimal

import pytest
from django.test import Client
from openpyxl import load_workbook

from apps.bpp.services.plan import service as plan
from apps.bpp.services.requests import printing as request_printing
from apps.bpp.services.requests import read
from apps.bpp.services.core import printing
from apps.bpp.services.requests import requests as service
from apps.bpp.tests import stage2 as s
from htqweb.errors import DomainError

pytestmark = pytest.mark.django_db
BASE = "/api/bpp/v1"
GD, ADM = 910, 909


def _sheet(response) -> list[tuple]:
    assert response.status_code == 200, response.content
    assert response["Content-Type"].startswith("application/vnd.openxmlformats")
    return list(load_workbook(io.BytesIO(response.content)).active.iter_rows(values_only=True))


def _submitted(slug, sn, proj, amount=100):
    req = service.create_draft(sn, {**s.header(proj, s.metal()),
                                    "items": s.items((1, amount))})
    return service.submit(sn, req.id, expected_version=None)


def _setup(slug):
    s.user(s.SN), s.user(s.SN2)
    proj = s.project(members=[s.SN, s.SN2])
    s.approved_budget(slug, proj, {s.metal(): 10_000})
    s.request_route()
    return proj


# ── реестр заявок ──────────────────────────────────────────────────────

def test_request_export_is_the_registry_the_user_sees(company_context):
    slug = company_context["slug"]
    proj = _setup(slug)
    mine = _submitted(slug, s.actor(slug, s.SN, "bpp-sn"), proj, 1250)
    _submitted(slug, s.actor(slug, s.SN2, "bpp-sn"), proj, 700)
    s.grant(slug, GD, "bpp-gd")

    rows = _sheet(Client().get(f"{BASE}/requests?format=xlsx", **s.auth(slug, s.SN)))

    assert rows[0][:3] == ("Номер", "Статус", "Дата создания")
    assert rows[0][-1] == "Сейчас у"
    assert [row[0] for row in rows[1:]] == [mine.number]   # чужая заявка СН не видна
    assert rows[1][1] == "На согласовании"
    assert Decimal(str(rows[1][8])) == Decimal("1250")   # число, не текст «1 250,00»
    assert rows[1][-1].startswith("ТД")

    # ГД не согласует заявки (маршрут «ТД → ОД»), а видит все — по узлу
    # bpp.requests.all, а не как участник согласования.
    director = _sheet(Client().get(f"{BASE}/requests?format=xlsx", **s.auth(slug, GD)))
    assert len(director) == 3


def test_background_rebuild_keeps_the_rights_of_whoever_ordered(company_context):
    """Фоновая пересборка идёт без запроса — права берутся по ``user_id``."""
    slug = company_context["slug"]
    proj = _setup(slug)
    _submitted(slug, s.actor(slug, s.SN, "bpp-sn"), proj)
    other = _submitted(slug, s.actor(slug, s.SN2, "bpp-sn"), proj)
    kwargs = {"user_id": s.SN2, "company": slug, "is_superuser": False, "filters": {}}

    assert read.export_count(**kwargs) == 1
    assert [row["number"] for row in read.export_rows(**kwargs)] == [other.number]


# ── реестр бюджетов и план ─────────────────────────────────────────────

def test_budget_export_has_limits_as_numbers(company_context):
    slug = company_context["slug"]
    proj = s.project()
    budget = s.approved_budget(slug, proj, {s.metal(): 10_000, s.design(): 5_000})
    s.grant(slug, s.FD, "bpp-fd")

    rows = _sheet(Client().get(f"{BASE}/budgets?format=xlsx", **s.auth(slug, s.FD)))

    assert rows[0][:4] == ("Номер", "Проект", "Наименование проекта", "Статус")
    assert rows[1][0] == budget.number and rows[1][1] == "П-015"
    assert rows[1][3] == "Утверждён"
    assert Decimal(str(rows[1][6])) == Decimal("15000")


def test_plan_export_lists_own_positions(company_context):
    slug = company_context["slug"]
    proj = _setup(slug)
    sn = s.actor(slug, s.SN, "bpp-sn")
    req = _submitted(slug, sn, proj, 300)
    s.decide(req, s.TD, "approve"), s.decide(req, s.OD, "approve")

    rows = _sheet(Client().get(f"{BASE}/plan?format=xlsx", **s.auth(slug, s.SN)))

    assert rows[0][0] == "Позиция" and rows[0][-1] == "Исполнитель"
    assert [row[0] for row in rows[1:]] == [f"{req.number}-01"]
    assert Decimal(str(rows[1][9])) == Decimal("300")


# ── печать заявки ──────────────────────────────────────────────────────

def test_print_shows_items_and_decisions_already_made(company_context):
    slug = company_context["slug"]
    proj = _setup(slug)
    req = _submitted(slug, s.actor(slug, s.SN, "bpp-sn"), proj, 1250)
    s.decide(req, s.TD, "approve", "Цены проверены")

    html = printing.render_html(request_printing.TEMPLATE, request_printing.context(req))

    assert "Заявка на закупку" in html and req.number in html
    assert "Позиция 1" in html and "1 250,00 KZT" in html
    assert "Лист согласования" in html
    assert "Согласовано" in html and "Цены проверены" in html
    assert "Ожидает решения" not in html   # ОД ещё не решил — в лист не попадает


def test_print_of_a_foreign_request_is_404(company_context):
    slug = company_context["slug"]
    proj = _setup(slug)
    req = _submitted(slug, s.actor(slug, s.SN, "bpp-sn"), proj)
    s.grant(slug, s.SN2, "bpp-sn")

    response = Client().get(f"{BASE}/requests/{req.pk}/print", **s.auth(slug, s.SN2))

    assert response.status_code == 404


# ── узлы прав (B-4) ────────────────────────────────────────────────────

def test_plan_reassign_needs_its_own_node(company_context):
    slug = company_context["slug"]
    proj = _setup(slug)
    sn = s.actor(slug, s.SN, "bpp-sn")
    req = _submitted(slug, sn, proj)
    s.decide(req, s.TD, "approve"), s.decide(req, s.OD, "approve")
    item = str(req.items.get().id)

    with pytest.raises(DomainError) as exc:
        plan.reassign(s.actor(slug, s.FD, "bpp-fd"), [item], to_user_id=s.SN2)
    assert exc.value.status == 403

    moved = plan.reassign(s.actor(slug, ADM, "bpp-adm"), [item], to_user_id=s.SN2)
    assert moved == 1
