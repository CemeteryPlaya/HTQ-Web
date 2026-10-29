"""Санитайзер формул xlsx (T-27): строка с первым символом ``= + - @ \\t \\r``
пишется ТЕКСТОМ — без формулы у читателя файла и без порчи значения."""

from __future__ import annotations

import io
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from htqweb.xlsx_safe import FORMULA_PREFIXES, is_dangerous, safe_cell

DANGEROUS = ["=HYPERLINK(\"http://x\")", "+77010000001", "-2+3", "@SUM(A1)", "\tcmd", "\rcmd"]


@pytest.mark.parametrize("value", DANGEROUS)
def test_dangerous_prefix_is_written_as_text(value):
    assert is_dangerous(value)
    book = Workbook(write_only=True)
    sheet = book.create_sheet("Лист")
    sheet.append([safe_cell(sheet, value)])
    buffer = io.BytesIO()
    book.save(buffer)
    cell = load_workbook(io.BytesIO(buffer.getvalue())).active["A1"]
    assert cell.data_type == "s"
    # Значение не испорчено (XML при чтении нормализует «\r» в «\n»).
    assert cell.value == value.replace("\r", "\n")


@pytest.mark.parametrize("value", ["Швеллер 12П", "", None, 42, Decimal("10.50"), "ТОО «А»"])
def test_safe_values_are_untouched(value):
    assert not is_dangerous(value)
    book = Workbook(write_only=True)
    sheet = book.create_sheet("Лист")
    cell = safe_cell(sheet, value)
    assert cell.value == value


def test_prefixes_are_the_owasp_six():
    assert set(FORMULA_PREFIXES) == {"=", "+", "-", "@", "\t", "\r"}
